"""Accessions and taxonomy IDs in the report link to their NCBI pages."""

from openpyxl import load_workbook

from qpcr_assay_check.config import load_config
from qpcr_assay_check.pipeline import evaluate
from qpcr_assay_check.report.html import render_report
from qpcr_assay_check.report.ncbi_links import accession_url, linkify, taxon_link
from qpcr_assay_check.report.xlsx import write_workbook

from .test_report_exclusivity import _specificity_with_exclusivity
from .test_variants import ASSAY, mk_site


def test_accession_urls_follow_ncbi_record_types():
    assert accession_url("NC_045512.2") == "https://www.ncbi.nlm.nih.gov/nuccore/NC_045512.2"
    assert accession_url("OZ558241.1").endswith("/nuccore/OZ558241.1")
    assert accession_url("NZ_CP012345.1").endswith("/nuccore/NZ_CP012345.1")
    assert accession_url("GCF_000008725.1") == (
        "https://www.ncbi.nlm.nih.gov/datasets/genome/GCF_000008725.1/"
    )
    for not_an_accession in ("BLASTN 2.16.0", "12345.6", "v1.1.1", "CTG4.1", "NC_045512"):
        assert accession_url(not_an_accession) is None


def test_free_text_is_escaped_and_only_accessions_become_links():
    text = "<b>closest</b> on PQ535257.1 (RID B7VYB28U016, 3.7 °C), e.g. GCA_000000003.1"
    html = str(linkify(text))
    assert html.startswith("&lt;b&gt;closest&lt;/b&gt;")
    assert html.count("<a ") == 2 and 'rel="noopener noreferrer"' in html
    assert '/nuccore/PQ535257.1"' in html and "/datasets/genome/GCA_000000003.1/" in html
    assert "wwwtax.cgi?id=2697049" in str(taxon_link(2697049))


def test_report_and_workbook_link_accessions_and_taxids(tmp_path):
    site = mk_site("forward", tier="exclusivity", acc="PQ535257.1", org="Influenza A virus",
                   s_aln="AAAT", midline="||| ", mm=1, level="warning")  # fmt: skip
    spec = _specificity_with_exclusivity().model_copy(update={"sites": [site]})
    cfg = load_config()
    result = evaluate(ASSAY, cfg, specificity=spec)
    html = render_report(result, cfg)
    assert '<a href="https://www.ncbi.nlm.nih.gov/nuccore/PQ535257.1"' in html
    assert f'<a href="https://www.ncbi.nlm.nih.gov/nuccore/{ASSAY.target.accession}"' in html
    assert "wwwtax.cgi?id=2697049" in html  # the intended target's taxonomy ID
    write_workbook(result, tmp_path / "r.xlsx")
    ws = load_workbook(tmp_path / "r.xlsx")["Off-target sites"]
    cell = next(c for c in ws["C"] if c.value == "PQ535257.1")
    assert cell.hyperlink.target == "https://www.ncbi.nlm.nih.gov/nuccore/PQ535257.1"
