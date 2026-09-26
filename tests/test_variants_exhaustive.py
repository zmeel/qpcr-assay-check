"""Exhaustive variant analysis (v1.1.0): locate, store, collect with a budget, assess, report."""

from __future__ import annotations

import json
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from qpcr_assay_check.config import load_config
from qpcr_assay_check.errors import InputError
from qpcr_assay_check.ncbi.http import NcbiHttp
from qpcr_assay_check.ncbi.settings import Credentials
from qpcr_assay_check.oligo import iupac
from qpcr_assay_check.pipeline import evaluate
from qpcr_assay_check.report.html import render_report
from qpcr_assay_check.report.xlsx import write_workbook
from qpcr_assay_check.variants.datasets import DatasetsClient
from qpcr_assay_check.variants.exhaustive import reference_amplicon, run_exhaustive
from qpcr_assay_check.variants.locate import find_loci, find_masked, find_masked_by_context
from qpcr_assay_check.variants.store import RegionStore

from .conftest import CDC_N1_F as F
from .conftest import CDC_N1_P as P
from .conftest import CDC_N1_R as R
from .conftest import make_assay
from .fake_datasets import FakeAssembly, FakeDatasets
from .world import filler, mutate

NOW = datetime(2026, 9, 23, tzinfo=UTC)
# A SYNTHETIC amplicon: the CDC N1 oligos with random spacers. Not a real sequence.
AMP = F + filler(30, 11) + P + filler(30, 12) + iupac.reverse_complement(R)
F_VARIANT = mutate(F, [20])  # a 3'-terminal forward-primer mismatch


def genome(seed: int, amp: str = AMP, *, reverse: bool = False) -> dict[str, str]:
    seq = filler(3000, seed) + amp + filler(3000, seed + 1)
    return {f"CTG{seed}.1": iupac.reverse_complement(seq) if reverse else seq}


def assemblies() -> list[FakeAssembly]:
    variant = AMP.replace(F, F_VARIANT, 1)
    broken = filler(3000, 7) + AMP[:60]
    return [
        FakeAssembly("GCF_000000001.1", "2024-03-01", genome(1)),
        FakeAssembly("GCF_000000002.1", "2025-05-01", genome(2, reverse=True)),
        FakeAssembly("GCA_000000003.1", "2026-01-15", genome(3, variant)),
        FakeAssembly("GCA_000000004.1", "2026-02-01", {"CTG4.1": filler(6000, 4)}),  # absent
        FakeAssembly("GCA_000000005.1", "2026-03-01",
                     {"CTG5a.1": broken, "CTG5b.1": AMP[60:] + filler(3000, 8)}),
    ]  # fmt: skip


def setup(tmp_path, *, budget: int = 20000, fake: FakeDatasets | None = None, key=None):
    cfg = load_config()
    cfg.variants.max_assemblies_per_run = budget
    cfg.ncbi.datasets_batch_size = 2
    fake = fake or FakeDatasets(assemblies())
    http = NcbiHttp(cfg.ncbi, Credentials("lab@example.org", key), session=fake,
                    sleep=lambda s: None, jitter=lambda: 0.0)  # fmt: skip
    client = DatasetsClient(http, cfg.ncbi.datasets_url)
    assay = make_assay(reference_amplicon=AMP, target={"taxid": 813})
    return cfg, fake, client, assay


def run(tmp_path, cfg, client, assay):
    return run_exhaustive(assay, cfg, client, tmp_path / "cache", _no_fetch, now=NOW)


def _no_fetch(acc: str) -> str:
    raise AssertionError("the reference amplicon was given; nothing should be fetched")


# ------------------------------------------------------------------ locate
def test_the_amplicon_is_found_on_either_strand_and_despite_a_primer_mismatch():
    kw = {"seed_length": 16, "seed_step": 4, "flank": 50}
    (fwd,) = find_loci(genome(1), AMP, **kw)
    (rev,) = find_loci(genome(2, reverse=True), AMP, **kw)
    (var,) = find_loci(genome(3, AMP.replace(F, F_VARIANT, 1)), AMP, **kw)
    for lc in (fwd, rev, var):
        assert not lc.truncated
    assert fwd.strand == "+" and rev.strand == "-"
    assert fwd.region[fwd.offset : fwd.offset + len(AMP)] == AMP
    assert rev.region[rev.offset : rev.offset + len(AMP)] == AMP  # read in the amplicon's sense
    assert var.region[var.offset : var.offset + len(F)] == F_VARIANT


def test_a_contig_break_is_flagged_and_an_absent_region_finds_nothing():
    kw = {"seed_length": 16, "seed_step": 4, "flank": 50}
    broken = {"a": filler(3000, 7) + AMP[:60], "b": AMP[60:] + filler(3000, 8)}
    assert all(lc.truncated for lc in find_loci(broken, AMP, **kw))
    assert find_loci({"x": filler(6000, 4)}, AMP, **kw) == []


# ------------------------------------------------------------------ the whole flow
def test_every_assembly_is_accounted_for_and_the_variant_is_counted(tmp_path):
    cfg, fake, client, assay = setup(tmp_path)
    res = run(tmp_path, cfg, client, assay)
    c = res.coverage
    assert (c.listed_total, c.assessed_total, c.processed_this_run) == (5, 5, 5)
    assert c.complete and (c.found, c.not_found, c.contig_break) == (3, 1, 1)
    assert c.not_found_examples == ["GCA_000000004.1"]
    assert {(y.year, y.listed, y.assessed) for y in c.years} == {
        (2026, 3, 3), (2025, 1, 1), (2024, 1, 1)
    }  # fmt: skip
    fwd = [s for s in res.sites if s.role == "forward"]
    assert sorted(s.n_mismatch for s in fwd) == [0, 0, 1]
    variant = next(s for s in fwd if s.n_mismatch)
    assert variant.accession == "GCA_000000003.1" and variant.terminal_defect
    rev_site = next(
        s for s in res.sites if s.accession == "GCF_000000002.1" and s.role == "forward"
    )
    assert rev_site.orientation == "-" and rev_site.n_mismatch == 0  # genome stored reverse


def test_one_copy_per_genbank_refseq_pair_and_the_api_key_header_are_requested(tmp_path):
    cfg, fake, client, assay = setup(tmp_path, key="secret-key")
    run(tmp_path, cfg, client, assay)
    reports = [c for c in fake.calls if c["url"].endswith("/dataset_report")]
    assert all(c["params"]["filters.exclude_paired_reports"] == "true" for c in reports)
    assert all(c["params"]["filters.assembly_version"] == "current" for c in reports)
    assert all(c["headers"].get("api-key") == "secret-key" for c in fake.calls)
    assert all("email" not in c["params"] for c in fake.calls)


def test_a_budget_processes_the_newest_first_and_the_next_run_resumes(tmp_path):
    cfg, fake, client, assay = setup(tmp_path, budget=3)
    first_result = run(tmp_path, cfg, client, assay)
    first = first_result.coverage
    assert "2025: 1 assembly listed, 0 assessed so far" in " ".join(
        first_result.inclusivity.rationale
    )
    assert "BLAST" not in " ".join(first_result.inclusivity.rationale)
    assert (first.processed_this_run, first.assessed_total, first.complete) == (3, 3, False)
    assert {y.year: y.assessed for y in first.years} == {2026: 3, 2025: 0, 2024: 0}
    n_downloads = len(fake.downloads)
    second = run(tmp_path, cfg, client, assay).coverage
    assert (second.processed_this_run, second.assessed_total, second.complete) == (2, 5, True)
    third = run(tmp_path, cfg, client, assay).coverage
    assert third.processed_this_run == 0
    # one download per release-year window (2025, 2024); nothing is ever downloaded twice
    assert len(fake.downloads) == n_downloads + 2


def test_a_failed_download_is_retried_on_the_next_run(tmp_path):
    fake = FakeDatasets(assemblies(), missing_from_download={"GCF_000000001.1"})
    cfg, fake, client, assay = setup(tmp_path, fake=fake)
    c = run(tmp_path, cfg, client, assay).coverage
    assert (c.download_failed_this_run, c.assessed_total, c.complete) == (1, 4, False)
    fake.missing_from_download.clear()
    assert run(tmp_path, cfg, client, assay).coverage.complete


def test_an_assembly_that_keeps_failing_is_left_out_without_blocking_completion(tmp_path):
    """Live Neisseria runs (user, 2026-09-26): the same 39 downloads failed on every run, so the
    analysis stayed 'incomplete' with 'run again to continue' forever."""
    fake = FakeDatasets(assemblies(), missing_from_download={"GCF_000000001.1"})
    cfg, fake, client, assay = setup(tmp_path, fake=fake)
    first = run(tmp_path, cfg, client, assay).coverage
    assert (first.unavailable, first.complete) == (0, False)  # one failure: try again
    second = run(tmp_path, cfg, client, assay).coverage
    assert second.download_failed_this_run == 1  # still tried
    assert (second.unavailable, second.complete) == (1, True)
    assert second.unavailable_examples == ["GCF_000000001.1"]
    res = run(tmp_path, cfg, client, assay)
    result = evaluate(assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
                      release_dates=res.release_dates, inclusivity=res.inclusivity,
                      specificity=_empty_specificity())  # fmt: skip
    text = " ".join(result.findings) if hasattr(result, "findings") else result.model_dump_json()
    assert "could not be downloaded after repeated attempts" in text
    assert "Run again to continue" not in text
    assert "could not be downloaded after repeated attempts" in render_report(result, cfg)
    fake.missing_from_download.clear()  # available again: stored, no longer unavailable
    third = run(tmp_path, cfg, client, assay).coverage
    assert (third.unavailable, third.assessed_total, third.complete) == (0, 5, True)


def test_the_region_store_survives_a_restart(tmp_path):
    cfg, fake, client, assay = setup(tmp_path)
    run(tmp_path, cfg, client, assay)
    (path,) = (tmp_path / "cache" / "variants").glob("813-*.jsonl")
    store = RegionStore(path)
    assert len(store.items) == 5 and "GCA_000000004.1" in store
    assert store.items["GCA_000000004.1"].status == "not_found"
    assert all(len(lc.region) < 400 for it in store.items.values() for lc in it.loci)  # no genomes


def test_no_assemblies_is_an_input_error_not_an_empty_pass(tmp_path):
    cfg, fake, client, assay = setup(tmp_path, fake=FakeDatasets([]))
    with pytest.raises(InputError, match="no genome assemblies"):
        run(tmp_path, cfg, client, assay)


def test_the_reference_amplicon_can_be_cut_from_the_target_accession():
    seq = filler(500, 1) + AMP + filler(500, 2)
    assay = make_assay(target={"taxid": 813, "accession": "NC_000117.1"})
    amp, source = reference_amplicon(
        assay, lambda acc: f">{acc}\n{iupac.reverse_complement(seq)}\n"
    )
    assert amp == AMP and "NC_000117.1" in source
    with pytest.raises(InputError, match="reference_amplicon"):
        reference_amplicon(assay, lambda acc: f">{acc}\n{filler(900, 3)}\n")


# ------------------------------------------------------------------ report
def test_the_report_and_workbook_show_coverage_and_first_release_dates(tmp_path):
    cfg, fake, client, assay = setup(tmp_path, budget=3)
    res = run(tmp_path, cfg, client, assay)
    result = evaluate(
        assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
        release_dates=res.release_dates, inclusivity=res.inclusivity,
        specificity=_empty_specificity(),
    )  # fmt: skip
    vs = result.variant_summary
    assert vs.source == "datasets" and not vs.target_list_full
    fwd = next(o for o in vs.oligos if o.role == "forward")
    variant = next(r for r in fwd.rows if r.n_mismatch)
    assert variant.first_seen == variant.last_seen == "2026-01-15"
    assert any("3 of 5 genome assemblies" in line for line in result.overall.rationale)
    html = render_report(result, cfg)
    assert "Scope: every genome assembly assessed so far." in html and "incomplete" in html
    assert "First / last release" in html
    write_workbook(result, tmp_path / "r.xlsx")
    rows = list(load_workbook(tmp_path / "r.xlsx")["Variant coverage"].iter_rows(values_only=True))
    assert ("Total", 5, 3) in rows


def _empty_specificity():
    from .test_report_exclusivity import _specificity_with_exclusivity

    return _specificity_with_exclusivity()


# ------------------------------------------------------------------ plasmid-borne targets
CHROM = "chromosome, complete genome"
PLASMID = "plasmid pCT, complete sequence"


def plasmid_assemblies() -> list[FakeAssembly]:
    def asm(acc, date, plasmid_seq):
        contigs = {f"{acc}_chr": filler(4000, hash(acc) % 1000)}
        desc = {f"{acc}_chr": f"Chlamydia trachomatis {CHROM}"}
        if plasmid_seq is not None:
            contigs[f"{acc}_pl"] = plasmid_seq
            desc[f"{acc}_pl"] = f"Chlamydia trachomatis {PLASMID}"
        return FakeAssembly(acc, date, contigs, descriptions=desc)

    return [
        asm("GCF_100.1", "2025-01-01", filler(500, 1) + AMP + filler(500, 2)),
        asm("GCF_101.1", "2025-02-01", filler(500, 3) + AMP + filler(500, 4)),
        asm("GCF_102.1", "2025-03-01", None),  # chromosome only: says nothing about the strain
        asm("GCF_103.1", "2026-01-01", filler(1200, 5)),  # plasmid without the region: review
    ]


def test_a_plasmid_target_separates_missing_plasmids_from_a_missing_region(tmp_path):
    cfg, fake, client, assay = setup(tmp_path, fake=FakeDatasets(plasmid_assemblies()))
    res = run(tmp_path, cfg, client, assay)
    c = res.coverage
    assert c.target_on_plasmid is True and c.found == 2 and c.not_found == 2
    assert (c.not_found_without_plasmid, c.not_found_with_plasmid) == (1, 1)
    assert c.not_found_with_plasmid_examples == ["GCF_103.1"]
    assert any(PLASMID in h for h in c.plasmid_header_examples)
    assert any("plasmid sequence but not the target region" in r for r in res.inclusivity.rationale)
    result = evaluate(
        assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
        release_dates=res.release_dates, inclusivity=res.inclusivity,
        specificity=_empty_specificity(),
    )  # fmt: skip
    assert any("GCF_103.1" in line and "nvCT" in line for line in result.overall.rationale)
    html = render_report(result, cfg)
    assert "labelled as a plasmid but not the target region" in html and "GCF_103.1" in html
    assert "Inclusivity across the intended target (all genome assemblies)" in html
    write_workbook(result, tmp_path / "r.xlsx")
    rows = list(load_workbook(tmp_path / "r.xlsx")["Variant coverage"].iter_rows(values_only=True))
    assert any(
        str(r[0]).strip().startswith("...plasmid sequence present") and r[1] == 1 for r in rows
    )


def test_not_found_entries_stored_before_the_plasmid_check_are_scanned_again(tmp_path):
    cfg, fake, client, assay = setup(tmp_path, fake=FakeDatasets(plasmid_assemblies()))
    run(tmp_path, cfg, client, assay)
    (path,) = (tmp_path / "cache" / "variants").glob("813-*.jsonl")
    # rewrite the store as an older version would have: no plasmid counts on 'not found' lines
    lines = []
    for line in path.read_text().splitlines():
        item = json.loads(line)
        if item["status"] == "not_found":
            item.pop("plasmid_contigs"), item.pop("n_contigs"), item.pop("plasmid_examples")
        lines.append(json.dumps(item))
    path.write_text("\n".join(lines) + "\n")
    before = len(fake.downloads)
    c = run(tmp_path, cfg, client, assay).coverage
    assert c.processed_this_run == 2 and len(fake.downloads) > before  # only the two 'not found'
    assert (c.not_found_without_plasmid, c.not_found_with_plasmid) == (1, 1)


def test_variant_tables_describe_the_match_in_words(tmp_path):
    cfg, fake, client, assay = setup(tmp_path)
    res = run(tmp_path, cfg, client, assay)
    result = evaluate(
        assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
        release_dates=res.release_dates, inclusivity=res.inclusivity,
        specificity=_empty_specificity(),
    )  # fmt: skip
    html = render_report(result, cfg)
    # per oligo (advisor 2026-09-25): perfect as one line, other variants by their changes
    i = html.index("<h3>Variants per oligo")
    per_oligo = html[i : html.index("<h2>", i)]
    assert "Perfect in <strong>" in per_oligo
    # the 3'-terminal forward mismatch is seen in one record: one row for its class
    assert "other likely failure variant, each seen in 1 record" in per_oligo
    assert html.index("<h3>Whole fragment") < i  # the whole fragment first
    # whole fragment (advisor layout 2026-09-25): the 3'-terminal forward mismatch needs attention
    i = html.index("<h3>Whole fragment")
    frag = html[i : html.index("<h2>", i)]
    assert "Needs attention" in frag and "likely failure" in frag and "Detectable (" in frag
    assert frag.index("likely failure") < frag.index("Detectable (")
    # the frequency of a non-perfect site variant sits next to its class
    assert 'title="this forward site variant, whatever the other sites are"' in frag


def test_found_entries_without_plasmid_info_are_rescanned_so_the_split_can_be_shown(tmp_path):
    """Live finding: found entries from the first v1.1.0 run had no plasmid info, which hid it."""
    cfg, fake, client, assay = setup(tmp_path, fake=FakeDatasets(plasmid_assemblies()))
    run(tmp_path, cfg, client, assay)
    (path,) = (tmp_path / "cache" / "variants").glob("813-*.jsonl")
    lines = []
    for line in path.read_text().splitlines():
        item = json.loads(line)
        for key in ("plasmid_contigs", "n_contigs", "plasmid_examples"):
            item.pop(key)
        for lc in item["loci"]:
            lc.pop("on_plasmid")
        lines.append(json.dumps(item))
    path.write_text("\n".join(lines) + "\n")
    c = run(tmp_path, cfg, client, assay).coverage
    assert c.processed_this_run == 4 and c.target_on_plasmid is True
    assert (c.not_found_without_plasmid, c.not_found_with_plasmid) == (1, 1)
    assert run(tmp_path, cfg, client, assay).coverage.processed_this_run == 0  # only once


def test_exhaustive_inclusivity_explains_assemblies_without_the_region(tmp_path):
    cfg, fake, client, assay = setup(tmp_path)
    res = run(tmp_path, cfg, client, assay)
    assert any(
        "2 of 5 assessed assemblies are not in the counts above" in r
        for r in res.inclusivity.rationale
    )
    result = evaluate(
        assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
        release_dates=res.release_dates, inclusivity=res.inclusivity,
        specificity=_empty_specificity(),
    )  # fmt: skip
    html = render_report(result, cfg)
    assert "<th>Assemblies</th><th>With region</th>" in html


# ------------------------------------------------------------------ history: new variants
def _evaluate(assay, cfg, res, now, previous=None):
    return evaluate(
        assay, cfg, now=now, target_sites=res.sites, variant_coverage=res.coverage,
        release_dates=res.release_dates, inclusivity=res.inclusivity,
        specificity=_empty_specificity(), previous_run=previous,
    )  # fmt: skip


def test_the_history_flags_emerging_and_newly_assessed_variants(tmp_path):
    base = [
        FakeAssembly("GCF_000000001.1", "2024-03-01", genome(1)),
        FakeAssembly("GCF_000000002.1", "2025-05-01", genome(2)),
    ]
    cfg, _, client, assay = setup(tmp_path, fake=FakeDatasets(list(base)))
    first = _evaluate(assay, cfg, run(tmp_path, cfg, client, assay),
                      datetime(2026, 9, 1, tzinfo=UTC))  # fmt: skip

    probe_variant = AMP.replace(P, mutate(P, [3]), 1)
    later = base + [
        FakeAssembly("GCA_000000003.1", "2026-09-10", genome(3, AMP.replace(F, F_VARIANT, 1))),
        FakeAssembly("GCA_000000009.1", "2020-06-01", genome(9, probe_variant)),
    ]
    cfg, _, client, assay = setup(tmp_path, fake=FakeDatasets(later))
    second = _evaluate(assay, cfg, run(tmp_path, cfg, client, assay), NOW, previous=first)
    h = second.history
    assert h.variants_compared and {v.kind for v in h.variant_changes} == {"new"}
    fwd = next(v for v in h.variant_changes if v.role == "forward")
    probe = next(v for v in h.variant_changes if v.role == "probe")
    assert fwd.newly_released is True and fwd.concern  # a 3'-terminal primer mismatch
    assert probe.newly_released is False and not probe.concern  # older assembly, 1 mismatch
    assert h.verdict.value == "WARN"
    assert any("emerging" in line and "3'-end mismatch" in line for line in h.rationale)
    html = render_report(second, cfg)
    assert "Oligo sequence variants since the previous run" in html
    assert "emerging" in html and "newly assessed" in html


def test_variants_are_not_compared_across_different_sources(tmp_path):
    cfg, _, client, assay = setup(tmp_path)
    res = run(tmp_path, cfg, client, assay)
    first = _evaluate(assay, cfg, res, datetime(2026, 9, 1, tzinfo=UTC))
    first.variant_summary.source = "blast_hits"
    first.variant_summary.coverage = None
    second = _evaluate(assay, cfg, res, NOW, previous=first)
    assert not second.history.variants_compared and not second.history.variant_changes
    assert "Variants not compared" in second.history.variants_note


# ------------------------------------------------------------------ wholly masked (v1.1.1)
# Live: SARS-CoV-2 records with one run of 1,144 N over the whole N1 region (OZ558241.1) were
# reported as 'not found'. Synthetic stand-ins: the reference around the amplicon is LEFT/RIGHT.
LEFT, RIGHT = filler(1000, 50), filler(1000, 51)
REF = filler(2000, 60) + LEFT + AMP + RIGHT + filler(2000, 61)
HIDDEN = (
    filler(2000, 60) + LEFT[:600] + "N" * (400 + len(AMP) + 300) + RIGHT[300:] + filler(2000, 61)
)
KW = {"seed_length": 16, "flank": 50}


def test_a_region_wholly_hidden_by_n_is_placed_by_the_reference_around_it():
    assert find_loci({"c": HIDDEN}, AMP, seed_step=4, **KW) == []
    assert find_masked({"c": HIDDEN}, AMP, seed_step=4, **KW) == []  # no real base to match
    for seq, strand in ((HIDDEN, "+"), (iupac.reverse_complement(HIDDEN), "-")):
        (lc,) = find_masked_by_context({"c": seq}, AMP, LEFT, RIGHT, **KW)
        assert lc.strand == strand
        assert lc.region[lc.offset : lc.offset + len(AMP)] == "N" * len(AMP)
    (lc,) = find_masked_by_context({"c": HIDDEN}, AMP, LEFT, RIGHT, **KW)
    assert lc.start == 3000 - 50 - 20 + 1  # the amplicon starts after 3000 bases; flank + pad


def test_real_bases_where_the_amplicon_should_be_are_not_called_masked():
    divergent = LEFT + filler(len(AMP), 53) + RIGHT + "N" * 200  # absent or too divergent
    assert find_masked_by_context({"c": divergent}, AMP, LEFT, RIGHT, **KW) == []
    assert find_masked_by_context({"c": HIDDEN}, AMP, "", "", **KW) == []  # no context: no call


def _masked_setup(tmp_path):
    fake = FakeDatasets([
        FakeAssembly("GCF_000000001.1", "2024-03-01", genome(1)),
        FakeAssembly("GCA_000000009.1", "2026-04-01", {"CTG9.1": HIDDEN}),
        FakeAssembly("GCA_000000004.1", "2026-02-01", {"CTG4.1": filler(6000, 4)}),  # absent
    ])  # fmt: skip
    cfg, fake, client, _assay = setup(tmp_path, fake=fake)
    assay = make_assay(reference_amplicon=AMP, target={"taxid": 813, "accession": "NC_000117.1"})
    fetched: list[str] = []

    def fetch(acc: str) -> str:
        fetched.append(acc)
        return f">{acc}\n{REF}\n"

    def go():
        return run_exhaustive(assay, cfg, client, tmp_path / "cache", fetch, now=NOW)

    return go, fetched, fake


def test_a_wholly_masked_assembly_is_counted_as_hidden_by_n_not_as_not_found(tmp_path):
    go, fetched, _fake = _masked_setup(tmp_path)
    c = go().coverage
    assert (c.found, c.masked, c.not_found) == (1, 1, 1)
    assert c.masked_examples == ["GCA_000000009.1"] and c.not_found_examples == ["GCA_000000004.1"]
    assert fetched == ["NC_000117.1"]  # the reference, once, only because something was not found
    assert go().coverage.masked == 1 and fetched == ["NC_000117.1"]  # nothing rescanned or fetched


def test_not_found_entries_stored_before_v1_1_1_are_checked_once_for_a_masked_region(tmp_path):
    go, fetched, fake = _masked_setup(tmp_path)
    go()
    (path,) = (tmp_path / "cache" / "variants").glob("813-*.jsonl")
    # as v1.1.0 stored them: the masked one as 'not found', neither with context_checked
    lines = []
    for line in path.read_text().splitlines():
        item = json.loads(line)
        item.pop("context_checked", None)
        if item["status"] == "masked":
            item.update(status="not_found", loci=[], n_loci=0)
        lines.append(json.dumps(item))
    path.write_text("\n".join(lines) + "\n")
    for f in path.parent.glob("*.context.json"):
        f.unlink()  # and without the cached reference context
    c = go().coverage
    assert c.processed_this_run == 2 and (c.masked, c.not_found) == (1, 1)
    assert go().coverage.processed_this_run == 0  # checked once, not every run


def test_inclusivity_has_a_whole_fragment_row_per_year(tmp_path):
    """User 2026-09-25: next to the per-oligo tables, one table with the genome outcome of the
    three sites together; it agrees with the copy coverage."""
    cfg, fake, client, assay = setup(tmp_path)
    res = run(tmp_path, cfg, client, assay)
    years = res.inclusivity.fragment_years
    assert years and all(
        f.detectable + f.at_risk + f.likely_failure + f.undetermined == f.with_region for f in years
    )
    cc = res.coverage.copies
    assert sum(f.detectable for f in years) == cc.with_detectable_copy
    assert sum(f.undetermined for f in years) == cc.undetermined
    result = evaluate(
        assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
        release_dates=res.release_dates, inclusivity=res.inclusivity,
        specificity=_empty_specificity(),
    )  # fmt: skip
    html = render_report(result, cfg)
    i = html.index("<h2>Inclusivity across")
    assert html.index("<h3>Whole fragment (forward + probe + reverse combined)</h3>", i) < (
        html.index("<h3>Forward</h3>", i)
    )
    # the window row uses the verdict's base (undetermined left out), so the numbers agree
    assert "Verdict window " in html and "verdict window</span>" in html
    assert "the same base as the verdict" in html


def test_the_inclusivity_verdict_uses_the_whole_fragment_over_recent_years():
    """Advisor 2026-09-26 (built on the user's request): whole-fragment outcome pooled over the
    last 3 complete years plus the current one; too few genomes = INCOMPLETE; a single large
    year below the FAIL limit gives at least WARN; small years do not decide."""
    from qpcr_assay_check.inclusivity.models import FragmentYear
    from qpcr_assay_check.variants.exhaustive import fragment_verdict
    from qpcr_assay_check.verdict import Verdict

    rules = load_config().inclusivity

    def year(y, det, risk=0, fail=0, undet=0):
        return FragmentYear(year=y, with_region=det + risk + fail + undet, detectable=det,
                            at_risk=risk, likely_failure=fail, undetermined=undet)  # fmt: skip

    good = [year(2014, 1, fail=4)] + [year(y, 98, risk=2) for y in range(2023, 2027)]
    verdict, lines = fragment_verdict(good, rules)
    assert verdict is Verdict.PASS and "2023-2026" in lines[0] and "98.0% detectable" in lines[0]
    # the old small year 2014 (1 of 5) does not decide; an undetermined genome is not counted
    warn = [year(y, 90, risk=8, fail=2, undet=5) for y in range(2023, 2027)]
    verdict, lines = fragment_verdict(warn, rules)
    assert verdict is Verdict.WARN and "90.0% detectable" in lines[0]
    assert "98.0% including at risk" in lines[0] and "not counted: 20" in lines[0]
    assert fragment_verdict([year(2026, 50, fail=50)], rules)[0] is Verdict.FAIL
    few = fragment_verdict([year(2026, 20)], rules)
    assert few[0] is Verdict.INCOMPLETE and "Too few recent genomes" in few[1][0]
    # one large year drops below the FAIL limit while the window passes: WARN
    drop = [year(y, 400) for y in range(2023, 2026)] + [year(2026, 20, fail=30)]
    verdict, lines = fragment_verdict(drop, rules)
    assert verdict is Verdict.WARN and any("Release year 2026 on its own" in x for x in lines)
