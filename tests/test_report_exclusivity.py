"""Exclusivity and taxonomy-breakdown sections of report.html and results.xlsx."""

from __future__ import annotations

from datetime import UTC, datetime

from openpyxl import load_workbook

from qpcr_assay_check.config import load_config
from qpcr_assay_check.pipeline import evaluate
from qpcr_assay_check.report.html import render_report
from qpcr_assay_check.report.xlsx import write_workbook
from qpcr_assay_check.search.orchestrate import SearchOutcome, SearchRecord
from qpcr_assay_check.specificity.models import SpecificityResult, TierCount
from qpcr_assay_check.taxonomy.plan import OrganismListResolution
from qpcr_assay_check.taxonomy.resolve import Resolution
from qpcr_assay_check.taxonomy.rollup import TaxonCount
from qpcr_assay_check.verdict import Verdict

from .test_exclusivity import CT, NG, amplicon, site

NOW = datetime(2026, 9, 20, 12, 0, 0, tzinfo=UTC)


def _specificity_with_exclusivity() -> SpecificityResult:
    sites = [site(CT, "forward", "critical", site_id="S1")]
    amps = [amplicon(CT, "likely_detected")]
    return SpecificityResult(
        verdict=Verdict.FAIL,
        verdict_sites=Verdict.FAIL,
        verdict_amplicons=Verdict.FAIL,
        scope="Off-target tiers assessed: exclusivity.",
        rationale=["Specificity: 1 critical primer site."],
        findings=[],
        counts=[
            TierCount(
                tier="exclusivity", query="forward", blast_hits=1, hsps_relevant=1,
                sites_assessed=1,
            )
        ],  # fmt: skip
        n_sites={"critical": 1, "warning": 0, "minor": 0},
        sites=sites,
        amplicons=amps,
        limitations=["Test limitation."],
        parameters={
            "alignment": {"match": 1, "mismatch": -3, "gap_open": 5, "gap_extend": 2},
            "primer_site": {}, "probe_site": {}, "severity": {},
            "window_padding_nt": 10, "max_amplicon_size": 2000, "min_identical_bases": 14,
            "blast": {"word_size": 7}, "databases": ["core_nt"],
            "blast_versions": ["BLASTN 2.17.0+"],
        },
    )  # fmt: skip


def _resolution() -> OrganismListResolution:
    return OrganismListResolution(
        resolutions=[
            Resolution(name="Chlamydia trachomatis", status="resolved", taxid=CT),
            Resolution(name="Neisseria gonorrhoeae", status="resolved", taxid=NG),
            Resolution(name="Mycoplasma pneumoniae", status="unresolved"),
        ]
    )


def _breakdown() -> list[TaxonCount]:
    return [
        TaxonCount(
            taxid=CT, scientific_name="Chlamydia trachomatis", species="Chlamydia trachomatis",
            genus="Chlamydia", family="Chlamydiaceae", n_sites=1,
        )
    ]  # fmt: skip


def _search_outcome() -> SearchOutcome:
    return SearchOutcome(
        tool={"name": "qpcr-assay-check", "version": "0"},
        generated_at=NOW.isoformat(),
        inputs_hash="x",
        oligos={},
        parameters={},
        searches=[
            SearchRecord(
                tier="exclusivity", label="exclusivity", taxids=[CT, NG], entrez_query=None,
                key="k", rid=None, state="done", blast_version=None, database=None,
                n_hits={}, saturation=[], restriction=None,
            )
        ],
    )  # fmt: skip


def _result(n1):
    cfg = load_config()
    specificity = _specificity_with_exclusivity()
    return evaluate(
        n1, cfg, now=NOW, specificity=specificity, organism_resolution=_resolution(),
        taxonomy_breakdown=_breakdown(), search_outcome=_search_outcome(),
    )  # fmt: skip


def test_the_exclusivity_table_shows_hit_and_zero_hit_and_unresolved_rows(n1):
    result = _result(n1)
    html = render_report(result, load_config())
    assert "Exclusivity against the clinical organism list" in html
    assert "2 of 3 organism-list name(s) resolved" in html
    assert "Chlamydia trachomatis" in html and "Neisseria gonorrhoeae" in html
    assert "Mycoplasma pneumoniae" in html and "unresolved" in html
    assert "likely_detected" in html


def test_off_target_sites_are_shown_per_species(n1):
    """The per-taxid breakdown (175 rows live) became one row per tier and species; the full
    breakdown with genus and family stays in the workbook."""
    html = render_report(_result(n1), load_config())
    assert "Off-target sites per species" in html
    assert "Chlamydia trachomatis" in html and "Chlamydiaceae" not in html


def test_no_exclusivity_or_breakdown_when_not_computed(n1):
    cfg = load_config()
    html = render_report(evaluate(n1, cfg, qc_only=True, now=NOW), cfg)
    assert "<h2>Exclusivity against the clinical organism list</h2>" not in html
    assert "<h2>Off-target sites per species</h2>" not in html


def test_workbook_gets_exclusivity_and_taxonomy_sheets(n1, tmp_path):
    result = _result(n1)
    path = tmp_path / "r.xlsx"
    write_workbook(result, path)
    wb = load_workbook(path)
    assert "Exclusivity" in wb.sheetnames and "Taxonomy breakdown" in wb.sheetnames
    rows = {row[0].value: row for row in wb["Exclusivity"].iter_rows(min_row=2)}
    assert rows["Chlamydia trachomatis"][3].value == 1  # Sites column
    assert rows["Mycoplasma pneumoniae"][1].value == "unresolved"
    tax_rows = list(wb["Taxonomy breakdown"].iter_rows(min_row=2))
    assert tax_rows[0][1].value == "Chlamydia trachomatis"


def test_a_run_without_a_human_search_says_so_in_the_rationale_and_report(n1):
    result = _result(n1)  # the search outcome has an exclusivity tier only
    assert any("Human background" in line for line in result.overall.rationale)
    spec = next(s for s in result.sections if s.key == "specificity")
    assert "Human background was not searched" in spec.note
    assert "off-target binding to human DNA was not evaluated" in render_report(
        result, load_config()
    )


def test_a_run_with_a_human_search_does_not_flag_it(n1):
    outcome = _search_outcome()
    outcome.searches.append(
        outcome.searches[0].model_copy(update={"tier": "background", "taxids": [9606]})
    )
    result = evaluate(
        n1, load_config(), now=NOW, specificity=_specificity_with_exclusivity(),
        organism_resolution=_resolution(), search_outcome=outcome,
    )  # fmt: skip
    assert not any("Human background" in line for line in result.overall.rationale)


def test_qc_only_runs_do_not_mention_human_background(n1):
    result = evaluate(n1, load_config(), qc_only=True, now=NOW)
    assert not any("Human background" in line for line in result.overall.rationale)


# ------------------------------------------------ variant summary: full target hit list, underline
def _target_record(list_full: bool):
    from qpcr_assay_check.search.assess import QuerySaturation

    return SearchRecord(
        tier="target", label="target", taxids=[2697049], entrez_query=None, key="t", rid=None,
        state="done", blast_version=None, database=None, n_hits={}, restriction=None,
        saturation=[QuerySaturation(
            label="forward", n_hits=5000 if list_full else 10, hitlist_size=5000,
            list_full=list_full, weakest_identity=20, min_relevant_identity=14,
            saturated=list_full, note="",
        )],
    )  # fmt: skip


def _variant_result(n1, *, list_full: bool, n_rare: int = 0):
    from .test_variants import mk_site

    sites = [mk_site(role, acc=f"A{i}.1") for i in range(2000) for role in ("forward", "probe")]
    sites += [mk_site("forward", acc=f"R{i}.1", s_aln="AAAT", mm=1) for i in range(n_rare)]
    outcome = _search_outcome()
    outcome.searches.append(_target_record(list_full))
    return evaluate(
        n1, load_config(), now=NOW, specificity=_specificity_with_exclusivity(),
        organism_resolution=_resolution(), search_outcome=outcome, target_sites=sites,
    )  # fmt: skip


def test_a_full_target_hit_list_marks_the_variant_tables_as_biased(n1, tmp_path):
    result = _variant_result(n1, list_full=True)
    assert result.variant_summary.target_list_full
    html = render_report(result, load_config())
    assert "Biased toward perfect matches." in html
    assert "separate selection of records" in html
    write_workbook(result, tmp_path / "r.xlsx")
    summary = {
        r[0].value: r[1].value for r in load_workbook(tmp_path / "r.xlsx")["Summary"].iter_rows()
    }
    assert "biased toward perfect matches" in summary["Variant tables"]


def test_a_target_hit_list_with_room_to_spare_is_not_marked_biased(n1):
    result = _variant_result(n1, list_full=False)
    assert not result.variant_summary.target_list_full
    assert "Biased toward perfect matches." not in render_report(result, load_config())


def test_a_rare_variant_shows_as_below_0_1_percent_not_zero(n1):
    html = render_report(_variant_result(n1, list_full=False, n_rare=1), load_config())
    i = html.index("<h3>Variants per oligo")
    per_oligo = html[i : html.index("<h2>", i)]
    assert '<span class="meta">&lt;0.1%</span>' in per_oligo and ">0.0%<" not in per_oligo


def test_only_primers_get_the_3prime_underline():
    from qpcr_assay_check.report.html import _alignment_html, _seq_html

    from .test_variants import mk_site

    assert 'class="tail"' in _seq_html("ACGTACGTAC")
    assert 'class="tail"' not in _seq_html("ACGTACGTAC", 0)
    assert "tail" in _alignment_html(mk_site("forward"))
    assert "tail" not in _alignment_html(mk_site("probe"))


def test_with_sampled_hits_the_fragment_table_states_its_own_coverage(n1):
    """Advisor 2026-09-25: with blast_hits the fragment table holds fewer records than the
    per-oligo tables, and says so."""
    from .test_variants import mk_site

    roles = ("forward", "probe", "reverse")
    sites = [mk_site(role, acc=f"A{i}.1") for i in range(3) for role in roles]
    sites += [mk_site("forward", acc=f"B{i}.1") for i in range(4)]
    result = evaluate(n1, load_config(), now=NOW, target_sites=sites, qc_only=False,
                      specificity=_specificity_with_exclusivity())  # fmt: skip
    html = render_report(result, load_config())
    assert "This table covers the 3 records" in html and "forward 7" in html
