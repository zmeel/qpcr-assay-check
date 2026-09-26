"""v1.3.0 step 2: best-binding copy per genome, coverage per oligo and channel, homopolymers.

All sequences here are SYNTHETIC (random spacers around the CDC N1 oligos), not real genomes.
"""

from __future__ import annotations

from qpcr_assay_check.oligo import iupac
from qpcr_assay_check.variants.exhaustive import run_exhaustive

from .conftest import CDC_N1_F as F
from .conftest import CDC_N1_P as P
from .conftest import CDC_N1_R as R
from .conftest import make_assay
from .fake_datasets import FakeAssembly, FakeDatasets
from .test_variants_exhaustive import AMP, NOW, _no_fetch, setup
from .world import filler, mutate

RC_R = iupac.reverse_complement(R)
SPACER1 = AMP[len(F) : len(F) + 30]
P2 = mutate(P, [2, 3, 5])  # an alternative probe for a divergent lineage (synthetic)


def copies(seed: int, *amps: str) -> dict[str, str]:
    """One contig holding every given copy of the region, separated by random sequence."""
    parts = [filler(2000, seed)]
    for i, amp in enumerate(amps):
        parts += [amp, filler(2000, seed + 100 + i)]
    return {f"CTG{seed}.1": "".join(parts)}


def run(tmp_path, assemblies, **assay_extra):
    cfg, fake, client, _ = setup(tmp_path, fake=FakeDatasets(assemblies))
    for key, value in assay_extra.pop("variants", {}).items():
        setattr(cfg.variants, key, value)
    assay = make_assay(**({"reference_amplicon": AMP, "target": {"taxid": 813}} | assay_extra))
    return run_exhaustive(assay, cfg, client, tmp_path / "cache", _no_fetch, now=NOW)


def test_the_genome_is_judged_by_its_best_binding_copy_not_the_first_found(tmp_path):
    # copy A: a reverse-primer 3'-end mismatch (few seeds lost, so it is found most confidently)
    # copy B: perfect oligo sites but a heavily changed spacer (many seeds lost)
    bad_3prime = AMP.replace(RC_R, mutate(RC_R, [1]))
    changed_spacer = AMP.replace(SPACER1, mutate(SPACER1, [3, 7, 11, 15, 19, 23, 27]))
    res = run(tmp_path, [FakeAssembly("GCA_000000101.1", "2026-02-01",
                                      copies(1, bad_3prime, changed_spacer))])  # fmt: skip
    rev = next(s for s in res.sites if s.role == "reverse")
    assert (rev.n_mismatch, rev.n_gap) == (0, 0)  # the perfect copy counts
    c = res.coverage.copies
    assert (c.genomes, c.multi_copy, c.max_copies) == (1, 1, 2)
    assert c.best_copy_not_first == 1 and c.with_detectable_copy == 1 and c.escapes == 0


def test_a_genome_without_any_detectable_copy_is_an_escape(tmp_path):
    both_bad = AMP.replace(RC_R, mutate(RC_R, [1, 2]))
    res = run(tmp_path, [
        FakeAssembly("GCA_000000102.1", "2026-02-01", copies(2, AMP)),
        FakeAssembly("GCA_000000103.1", "2026-03-01", copies(3, both_bad, both_bad)),
    ])  # fmt: skip
    c = res.coverage.copies
    assert c.escapes == 1 and c.escape_examples == ["GCA_000000103.1"]
    assert c.role_none["reverse"] == 1 and c.role_none_examples["reverse"] == ["GCA_000000103.1"]
    assert c.role_none["forward"] == 0


def test_coverage_per_oligo_and_per_probe_channel(tmp_path):
    lineage2 = AMP.replace(P, P2)
    probes = [{"name": "P1", "sequence": P, "reporter": "FAM"},
              {"name": "P2", "sequence": P2, "reporter": "HEX"}]  # fmt: skip
    genomes = [
        FakeAssembly("GCA_000000104.1", "2026-02-01", copies(4, AMP)),
        FakeAssembly("GCA_000000105.1", "2026-03-01", copies(5, lineage2)),
        FakeAssembly("GCA_000000106.1", "2026-04-01", copies(6, AMP, lineage2)),
    ]
    c = run(tmp_path, genomes, probe=probes).coverage.copies
    rows = {r.name: (r.covered, r.only) for r in c.oligos if r.role == "probe"}
    # on each genome's best copy: 104 and 106 -> P1 (first best copy), 105 -> P2
    assert rows["P1"][0] >= 1 and rows["P2"][0] >= 1 and c.role_none["probe"] == 0
    assert {r.reporter: r.covered for r in c.channels} == {"FAM": rows["P1"][0],
                                                         "HEX": rows["P2"][0]}  # fmt: skip
    assert c.any_channel == 3 and c.probe_channels == "any" and c.escapes == 0


def test_all_channels_requires_every_reporter_to_detect(tmp_path):
    probes = [{"name": "P1", "sequence": P, "reporter": "FAM"},
              {"name": "P2", "sequence": P2, "reporter": "HEX"}]  # fmt: skip
    genomes = [FakeAssembly("GCA_000000107.1", "2026-02-01", copies(7, AMP))]
    res = run(tmp_path, genomes, probe=probes, variants={"probe_channels": "all"})
    probe = next(s for s in res.sites if s.role == "probe")
    assert probe.query == "P2" and probe.n_mismatch == 3  # the HEX channel does not detect
    assert res.coverage.copies.escapes == 1 and res.coverage.copies.all_channels == 0


def test_a_homopolymer_run_length_variant_is_aligned_as_a_bulge_and_labelled(tmp_path):
    longer_run = F.replace("AAAA", "AAAAA", 1)  # GACCCCAAAAA... : poly-A 4 -> 5
    res = run(tmp_path, [FakeAssembly("GCA_000000108.1", "2026-02-01",
                                      copies(8, AMP.replace(F, longer_run)))])  # fmt: skip
    fwd = next(s for s in res.sites if s.role == "forward")
    assert fwd.note == "poly-A run 4→5 (homopolymer length variant)"
    assert (fwd.n_mismatch, fwd.n_gap) == (0, 1) and fwd.mismatches_last5 == 0
    # R5b (advisor 2026-09-26): one base, run ending at -11, away from the last 3 nt: at risk
    assert (fwd.grade, fwd.grade_rule) == ("at_risk", "R5b")
    rl = res.coverage.copies.run_length
    assert (rl.genomes, rl.on_best_copy, rl.mixed, rl.decided_by_rule) == (1, 1, 0, 1)
    assert rl.variants == [("forward: poly-A run 4→5 (homopolymer length variant)", 1)]


def test_a_second_reference_finds_a_lineage_the_first_cannot(tmp_path):
    # a lineage without any exact 16-base stretch of the first reference (sites mutated every
    # 8 bases, other spacers), so only its own reference amplicon can find it
    other = filler(40, 71)
    lineage2 = (mutate(F, [5, 13]) + other[:30] + mutate(P, [6, 14, 22]) + other[10:40]
                + mutate(RC_R, [3, 11, 19]))  # fmt: skip
    genomes = [FakeAssembly("GCA_000000109.1", "2026-02-01", copies(9, lineage2))]
    res = run(tmp_path, genomes)  # one reference: the region is not found
    assert res.coverage.not_found == 1
    refs = [{"name": "L1", "sequence": AMP}, {"name": "L2", "sequence": lineage2}]
    res = run(tmp_path, genomes, reference_amplicon=None, reference_amplicons=refs)
    assert res.coverage.not_found == 0 and res.coverage.found == 1  # rescanned once, found


def run_report_result(tmp_path):
    """A full RunResult from a small SYNTHETIC multi-copy run (used by report tests)."""
    from qpcr_assay_check.config import load_config
    from qpcr_assay_check.pipeline import evaluate

    from .test_variants_exhaustive import _empty_specificity

    longer_run = F.replace("AAAA", "AAAAA", 1)
    both_bad = AMP.replace(RC_R, mutate(RC_R, [1, 2]))
    probes = [{"name": "P1", "sequence": P, "reporter": "FAM"},
              {"name": "P2", "sequence": P2, "reporter": "HEX"}]  # fmt: skip
    genomes = [
        FakeAssembly("GCA_000000110.1", "2026-02-01", copies(10, AMP.replace(F, longer_run))),
        FakeAssembly("GCA_000000111.1", "2026-03-01", copies(11, both_bad, AMP.replace(P, P2))),
        FakeAssembly("GCA_000000099.1", "2025-05-01", {"CTG99.1": filler(3000, 99)}),  # no region
    ]
    res = run(tmp_path, genomes, probe=probes)
    assay = make_assay(reference_amplicon=AMP, target={"taxid": 813}, probe=probes)
    cfg = load_config()
    return evaluate(assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
                    release_dates=res.release_dates, inclusivity=res.inclusivity,
                    specificity=_empty_specificity())  # fmt: skip


def test_the_report_and_workbook_show_copies_coverage_escapes_and_homopolymers(tmp_path):
    from openpyxl import load_workbook

    from qpcr_assay_check.config import load_config
    from qpcr_assay_check.report.html import render_report
    from qpcr_assay_check.report.xlsx import write_workbook

    result, cfg = run_report_result(tmp_path), load_config()
    html = render_report(result, cfg)
    assert "Copies, coverage per oligo, and escapes" in html and "none of them" in html
    assert "poly-A run 4→5 (homopolymer length variant)" in html
    assert "Probe channels" in html and "any channel" in html
    assert "not counted (strict)" in html and "if homopolymer bulges are tolerated" in html
    # graded classes: the columns and the explanation, although the first year (2017) is empty
    # (live 2026-09-25 they were hidden, as the template looked at the first year only)
    assert "<th>Detectable</th>" in html and "<h3>Mismatch classes</h3>" in html
    assert "Homopolymer length variants" in html and "copies disagree" in html
    # layout (user 2026-09-25): wider page; one-line alignments in the whole-fragment table
    assert "max-width: 96rem" in html and '<pre class="aln compact">' in html
    # alternatives numbered after the sequence, so the site lines stay aligned (user 2026-09-25)
    assert '<span class="seq" title="P1">' in html and "Probe: (1) P1, (2) P2." in html
    assert '<span class="meta">(2)</span></pre>' in html
    write_workbook(result, tmp_path / "r.xlsx")
    ws = load_workbook(tmp_path / "r.xlsx")["Copies and coverage"]
    assert any(c.value == "Escapes (no detectable copy)" for c in ws["A"])


def test_homopolymer_bulges_are_strict_by_default_and_both_counts_are_reported(tmp_path):
    longer_run = F.replace("AAAA", "AAAAA", 1)  # the only difference: poly-A 4 -> 5
    genomes = [FakeAssembly("GCA_000000112.1", "2026-02-01",
                            copies(12, AMP.replace(F, longer_run)))]  # fmt: skip
    strict = run(tmp_path, genomes).coverage.copies
    assert strict.homopolymer_bulges_detectable is False
    assert strict.with_detectable_copy == 0 and strict.escapes == 1
    assert (strict.with_detectable_copy_strict, strict.with_detectable_copy_bulges) == (0, 1)
    tolerant = run(tmp_path, genomes,
                   variants={"homopolymer_bulges_detectable": True}).coverage.copies  # fmt: skip
    assert tolerant.with_detectable_copy == 1 and tolerant.escapes == 0
    assert (tolerant.with_detectable_copy_strict, tolerant.with_detectable_copy_bulges) == (0, 1)


def test_a_bulge_with_a_mismatch_is_never_detectable():
    from qpcr_assay_check.variants.exhaustive import detectable

    site = type("S", (), {"n_gap": 1, "n_mismatch": 1, "mismatches_last5": 0, "note": "poly-A",
                          "grade": "indeterminate"})  # fmt: skip
    assert not detectable(site, True) and not detectable(site, False)  # type: ignore[arg-type]


def test_genomes_stored_with_the_old_copy_limit_are_downloaded_again_once(tmp_path, monkeypatch):
    from qpcr_assay_check.variants import store as store_mod

    def once():
        cfg, fake, client, _ = setup(tmp_path, fake=FakeDatasets(
            [FakeAssembly("GCA_000000113.1", "2026-02-01", copies(13, *[AMP] * 7))]))  # fmt: skip
        assay = make_assay(reference_amplicon=AMP, target={"taxid": 813})
        res = run_exhaustive(assay, cfg, client, tmp_path / "cache", _no_fetch, now=NOW)
        c = res.coverage.copies
        return c.max_copies, c.copies_capped, len(fake.downloads)

    monkeypatch.setattr(store_mod, "MAX_LOCI_KEPT", 5)  # a store written before v1.3.0
    assert once() == (5, 0, 1)
    monkeypatch.setattr(store_mod, "MAX_LOCI_KEPT", 20)
    assert once() == (7, 0, 1)  # downloaded again: every copy is stored now
    assert once() == (7, 0, 0)  # and only once


def test_a_genome_with_more_copies_than_kept_is_not_rescanned_every_run():
    from qpcr_assay_check.variants.store import MAX_LOCI_KEPT, StoredAssembly

    base = {"accession": "GCA_1.1", "release_date": "2026-01-01", "status": "found",
            "plasmid_contigs": 0}  # fmt: skip
    loci = [{"contig": "c", "strand": "+", "start": 1, "end": 2, "region": "A", "offset": 0,
             "n_seeds": 1, "truncated": False}]  # fmt: skip
    old = StoredAssembly(**base, n_loci=8, loci=loci * 5)
    many = StoredAssembly(**base, n_loci=MAX_LOCI_KEPT + 5, loci=loci * MAX_LOCI_KEPT)
    assert old.needs_rescan and old.copies_capped
    assert not many.needs_rescan and not many.copies_capped


def test_one_mismatch_in_an_mgb_probe_makes_a_genome_undetermined_not_an_escape(tmp_path):
    """User decision 2026-09-25 (the enterovirus MGB probe: 196 genomes had counted as escapes)."""
    probes = [{"name": "P", "sequence": P, "reporter": "FAM", "modifications": ["MGB"]}]
    one = AMP.replace(P, mutate(P, [8]))  # one mismatch in the probe site, away from its ends
    res = run(tmp_path, [FakeAssembly("GCA_000000120.1", "2026-02-01", copies(20, one))],
              probe=probes)  # fmt: skip
    c = res.coverage.copies
    assert (c.escapes, c.undetermined, c.with_detectable_copy) == (0, 1, 0)
    assert c.role_undetermined["probe"] == 1 and c.role_none["probe"] == 0
    (w,) = [w for o in res.inclusivity.oligos if o.role == "probe" for w in o.windows
            if w.sample_size]  # fmt: skip
    assert w.n_undetermined == 1


def test_copies_that_disagree_in_run_length_are_counted(tmp_path):
    """Advisor 2026-09-26: run length is a known sequencing/assembly error; a genome whose other
    copy reads the oligo's run length is counted as 'copies disagree'. SYNTHETIC genome."""
    longer_run = AMP.replace(F, F.replace("AAAA", "AAAAA", 1))
    other_bad = AMP.replace(RC_R, mutate(RC_R, [1, 2]))  # normal run, reverse site damaged
    res = run(tmp_path, [FakeAssembly("GCA_000000112.1", "2026-02-01",
                                      copies(12, longer_run, other_bad))])  # fmt: skip
    rl = res.coverage.copies.run_length
    assert (rl.genomes, rl.mixed) == (1, 1)
    assert sum(n for n, _t in rl.by_level.values()) == 1
