"""Several oligos per role (alternatives in one mix) and named oligos (v1.3.0 step 1)."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook
from pydantic import ValidationError

from qpcr_assay_check.config import load_config
from qpcr_assay_check.models import Assay
from qpcr_assay_check.oligo.qc import run_oligo_qc
from qpcr_assay_check.pipeline import evaluate
from qpcr_assay_check.report.html import render_report
from qpcr_assay_check.report.xlsx import write_workbook
from qpcr_assay_check.results import RunResult
from qpcr_assay_check.search.planner import build_queries
from qpcr_assay_check.variants.exhaustive import run_exhaustive

from .conftest import CDC_N1_F as F
from .conftest import CDC_N1_P as P
from .conftest import CDC_N1_R as R
from .conftest import make_assay
from .fake_datasets import FakeAssembly, FakeDatasets
from .test_variants_exhaustive import AMP, NOW, genome, setup
from .world import mutate

NG_EXAMPLE = (
    Path(__file__).parents[1] / "docs" / "examples" / "neisseria_gonorrhoeae_two_probes.yaml"
)
P2 = mutate(P, [2, 3, 5])  # a SYNTHETIC alternative probe for a divergent lineage


def two_probes(**extra: object) -> Assay:
    return make_assay(
        probe=[{"name": "P1", "sequence": P},
               {"name": "P2", "sequence": P2, "reporter": "HEX"}],
        **extra,
    )  # fmt: skip


# ------------------------------------------------------------------ assay file
def test_a_plain_sequence_still_works_and_is_named_after_its_role():
    a = make_assay()
    assert a.oligos == {"forward": F, "reverse": R, "probe": P}
    assert a.probe[0].reporter == "FAM" and not a.multi_oligo  # probe_reporter as default


def test_named_single_oligos_and_lists_are_accepted():
    a = two_probes(forward={"name": "F1", "sequence": F})
    assert list(a.oligos) == ["F1", "reverse", "P1", "P2"] and a.multi_oligo
    assert [(p.name, p.reporter) for p in a.probe] == [("P1", "FAM"), ("P2", "HEX")]
    assert a.role_of("P2") == "probe" and a.role_of("F1_v3") == "forward"


@pytest.mark.parametrize(
    ("probe", "message"),
    [
        ([{"sequence": P}, {"sequence": P2}], "give every oligo a name"),
        ([{"name": "P1", "sequence": P}, {"name": "P1", "sequence": P2}], "used twice"),
        ([{"name": "my probe", "sequence": P}], "no spaces"),
        ([{"name": "P_v2", "sequence": P}], "reserved for degenerate"),
        ([], "at least one oligo"),
    ],
)
def test_bad_oligo_lists_are_rejected_with_a_clear_message(probe, message):
    with pytest.raises(ValidationError, match=message):
        make_assay(probe=probe)


def test_a_saved_assay_validates_again_unchanged():
    a = two_probes(reference_amplicon=AMP)
    again = Assay.model_validate(a.model_dump(mode="json"))
    assert again == a and [r.name for r in again.reference_amplicons] == ["reference"]


def test_the_n_gonorrhoeae_example_is_valid():
    a = Assay.model_validate(yaml.safe_load(NG_EXAMPLE.read_text()))
    assert list(a.oligos) == ["NG-F", "NG-R", "NG-P1", "NG-P2"]
    assert [r.name for r in a.reference_amplicons] == ["NG-fragment-P1", "NG-fragment-P2"]
    assert all(p.modifications == ["MGB"] for p in a.probe)


# ------------------------------------------------------------------ QC
def test_qc_checks_every_oligo_and_every_pair_in_the_mix():
    cfg = load_config()
    qc = run_oligo_qc(two_probes(), cfg)
    subjects = {c.subject for c in qc.checks}
    assert {"forward", "reverse", "P1", "P2", "mix"} <= subjects
    labels = {s.label for s in qc.structures}
    assert "P1 / P2 dimer" in labels and "3' end of forward on P2" in labels
    spread = next(c for c in qc.checks if c.id == "mix.probe_tm_spread")
    assert spread.status.value == "INFO" and "P1, P2" in spread.name
    assert [o.name for o in qc.oligos] == ["forward", "reverse", "P1", "P2"]


def test_each_oligo_is_placed_in_the_reference_it_fits_best():
    cfg = load_config()
    lineage2 = AMP.replace(P, P2)
    a = two_probes(reference_amplicons=[{"name": "L1", "sequence": AMP},
                                        {"name": "L2", "sequence": lineage2}])  # fmt: skip
    checks = {c.id: c for c in run_oligo_qc(a, cfg).checks}
    assert checks["P1.site"].display.endswith("in L1, 0 mismatch(es)")
    assert checks["P2.site"].display.endswith("in L2, 0 mismatch(es)")
    assert checks["P2.within_amplicon"].status.value == "PASS"
    assert "amplicon.length.L1" in checks and "amplicon.length.L2" in checks


def test_an_alternative_that_fits_no_reference_warns_but_does_not_fail():
    cfg = load_config()
    far = mutate(P, [2, 3, 5, 8, 11, 14])  # beyond max_site_mismatches in the only reference
    a = make_assay(probe=[{"name": "P1", "sequence": P}, {"name": "Pfar", "sequence": far}],
                   reference_amplicon=AMP)  # fmt: skip
    checks = {c.id: c for c in run_oligo_qc(a, cfg).checks}
    assert checks["P1.site"].status.value == "PASS"
    assert (
        checks["Pfar.site"].status.value == "WARN"
        and "Another oligo" in checks["Pfar.site"].message
    )


# ------------------------------------------------------------------ searches
def test_blast_queries_are_labelled_by_oligo_name():
    cfg = load_config()
    q = build_queries(two_probes(forward={"name": "F1", "sequence": F[:-1] + "W"}), cfg)
    assert list(q) == ["F1_v1", "F1_v2", "reverse", "P1", "P2"]


# ------------------------------------------------------------------ variant analysis
def test_the_best_binding_alternative_counts_for_each_genome(tmp_path):
    fake = FakeDatasets([
        FakeAssembly("GCF_000000001.1", "2025-03-01", genome(1)),
        FakeAssembly("GCF_000000002.1", "2025-04-01", genome(2, AMP.replace(P, P2))),
    ])  # fmt: skip
    cfg, fake, client, _ = setup(tmp_path, fake=fake)
    assay = two_probes(reference_amplicon=AMP, target={"taxid": 813})
    res = run_exhaustive(assay, cfg, client, tmp_path / "cache", _no_fetch, now=NOW)
    probes = {s.accession: (s.query, s.n_mismatch) for s in res.sites if s.role == "probe"}
    assert probes == {"GCF_000000001.1": ("P1", 0), "GCF_000000002.1": ("P2", 0)}
    result = evaluate(assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
                      release_dates=res.release_dates, inclusivity=res.inclusivity,
                      specificity=_empty_specificity())  # fmt: skip
    probe_rows = next(o for o in result.variant_summary.oligos if o.role == "probe").rows
    assert sorted(r.oligo_name for r in probe_rows) == ["P1", "P2"]
    html = render_report(result, cfg)
    assert "P1, P2: for each record the best-binding one is shown" in html
    assert "<strong>P2</strong>:" in html and "Reaction mix" in html
    RunResult.model_validate_json(result.model_dump_json())  # a saved record loads again
    write_workbook(result, tmp_path / "r.xlsx")
    ws = load_workbook(tmp_path / "r.xlsx")["Oligo variants"]
    assert {c.value for c in ws["A"]} >= {"probe P1", "probe P2"}


def _no_fetch(acc: str) -> str:
    raise AssertionError("the reference amplicon was given; nothing should be fetched")


def _empty_specificity():
    from .test_report_exclusivity import _specificity_with_exclusivity

    return _specificity_with_exclusivity()


def test_a_role_counts_as_hit_on_its_target_when_any_of_its_oligos_is():
    """Live (NG two-probe run): named oligos were read as missing roles ("forward, probe,
    reverse" without a perfect hit) because the finding still parsed query labels."""
    from qpcr_assay_check.config import load_config
    from qpcr_assay_check.specificity.findings import build_findings

    roles = {"NG-F": "forward", "NG-R": "reverse", "NG-P1": "probe", "NG-P2": "probe"}
    common = {
        "sites": [], "amplicons": [], "site_by_id": {}, "counts": [], "saturated": [],
        "off_tiers_seen": [], "target_searched": True, "n_primer_only": 0,
        "n_fetch_failed": 0, "amplicons_truncated": False,
        "rules": load_config().specificity, "oligo_roles": roles,
    }  # fmt: skip
    hits = {"NG-F": 321, "NG-R": 297, "NG-P1": 321, "NG-P2": 0}
    (f,) = [x for x in build_findings(intended_target=hits, **common) if "Intended" in x.message]
    assert f.severity == "INFO" and "NG-P2 0" in f.message  # P1 covers the probe role
    no_probe = hits | {"NG-P1": 0}
    (f,) = [x for x in build_findings(intended_target=no_probe, **common)
            if "Intended" in x.message]  # fmt: skip
    assert f.severity == "WARN" and "no perfect full-length hit for probe (" in f.message
