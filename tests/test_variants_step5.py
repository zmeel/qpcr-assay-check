"""The analysis on the chain locator and store v2 (overhaul step 5), SYNTHETIC genomes."""

from qpcr_assay_check.oligo import iupac
from qpcr_assay_check.variants.exhaustive import run_exhaustive

from .conftest import make_assay
from .fake_datasets import FakeAssembly, FakeDatasets
from .test_variants_exhaustive import AMP, NOW, _no_fetch, run, setup
from .world import filler


def test_a_copy_with_a_long_insertion_is_judged_at_the_right_sites(tmp_path):
    """The old locator split such a copy and placed the far oligos 40 nt off (a false escape;
    measured live: 518 of 1,132 Legionella sites)."""
    longer = AMP[:30] + filler(40, 7) + AMP[30:]  # 40 nt in the spacer before the probe
    shorter = AMP[:24] + AMP[48:]  # 24 nt out of the same spacer
    one = {"c1": filler(900, 1) + longer + filler(900, 2)}
    two = {"c2": iupac.reverse_complement(filler(900, 3) + shorter + filler(900, 4))}
    fake = FakeDatasets(
        [
            FakeAssembly("GCF_000000001.1", "2025-01-01", one),
            FakeAssembly("GCF_000000002.1", "2025-02-01", two),
        ]
    )
    cfg, _f, client, assay = setup(tmp_path, fake=fake)
    res = run(tmp_path, cfg, client, assay)
    assert res.coverage.found == 2 and res.coverage.copies.escapes == 0
    assert all(s.n_mismatch == 0 and s.n_gap == 0 for s in res.sites)


def test_a_failed_batch_is_tried_again_in_halves(tmp_path):
    genomes = [FakeAssembly(f"GCF_00000000{i}.1", f"2025-0{i}-01",
                            {f"c{i}": filler(500, i) + AMP + filler(500, i + 10)})
               for i in range(1, 5)]  # fmt: skip
    fake = FakeDatasets(genomes, breaks_download={"GCF_000000003.1"})
    cfg, _f, client, assay = setup(tmp_path, fake=fake)
    cfg.ncbi.datasets_batch_size = 4
    res = run(tmp_path, cfg, client, assay)
    assert res.coverage.found == 3 and res.coverage.download_failed_this_run == 1
    store_failures = (tmp_path / "cache" / "genomes").glob("*.failures.json")
    (path,) = store_failures
    assert '"GCF_000000003.1"' in path.read_text()


def test_the_sequence_report_is_fetched_only_for_genomes_with_several_sequences(tmp_path):
    one = {"c1": filler(500, 1) + AMP + filler(500, 2)}
    two = {"c2": filler(500, 3) + AMP + filler(500, 4), "p2": filler(900, 5)}
    fake = FakeDatasets(
        [
            FakeAssembly("GCF_000000001.1", "2025-01-01", one),
            FakeAssembly("GCF_000000002.1", "2025-02-01", two, molecules={"p2": "Plasmid"}),
        ]
    )
    cfg, _f, client, assay = setup(tmp_path, fake=fake)
    res = run(tmp_path, cfg, client, assay)
    reports = [c["url"] for c in fake.calls if c["url"].endswith("/sequence_reports")]
    assert len(reports) == 1 and "GCF_000000002.1" in reports[0]
    assert res.coverage.found == 2


def test_nothing_is_downloaded_again_on_a_second_run(tmp_path):
    cfg, fake, client, assay = setup(tmp_path)
    first = run(tmp_path, cfg, client, assay)
    n = len(fake.downloads)
    second = run(tmp_path, cfg, client, assay)
    assert len(fake.downloads) == n
    assert [s.model_dump() for s in first.sites] == [s.model_dump() for s in second.sites]


# ------------------------------------------------------------------ 5c: per channel
def _swap(seq: str, positions: list[int]) -> str:
    alt = {"A": "C", "C": "G", "G": "T", "T": "A"}
    return "".join(alt[b] if i in positions else b for i, b in enumerate(seq))


def test_one_channel_counts_a_missing_locus_as_not_detected_only_in_complete_genomes(tmp_path):
    one = {"c1": filler(500, 1) + AMP + filler(500, 2)}
    fake = FakeDatasets([
        FakeAssembly("GCF_000000001.1", "2025-01-01", one),
        FakeAssembly("GCF_000000002.1", "2025-02-01", {"c2": filler(1500, 3)},
                     level="Complete Genome"),
        FakeAssembly("GCF_000000003.1", "2025-03-01", {"c3": filler(1500, 4)}),
    ])  # fmt: skip
    cfg, _f, client, assay = setup(tmp_path, fake=fake)
    (ch,) = run(tmp_path, cfg, client, assay).coverage.channel_results
    assert (ch.name, ch.target_genomes, ch.detected) == ("FAM", 3, 1)
    assert (ch.not_detected, ch.no_locus) == (1, 1)  # complete: a possible deletion; draft: no
    assert ch.not_detected_examples == ["GCF_000000002.1"] and ch.detected_percent == 50.0


def _two_channels(tmp_path):
    """Legionella-type: a genus channel and a species channel on one amplicon (SYNTHETIC)."""
    from qpcr_assay_check.models import Assay

    from .conftest import CDC_N1_F as F
    from .conftest import CDC_N1_P as P
    from .conftest import CDC_N1_R as R

    species = AMP[20:42]  # a second probe inside the same product
    assay = Assay.model_validate({
        "assay_name": "genus + species (synthetic)",
        "forward": {"name": "G-F", "sequence": F}, "reverse": {"name": "G-R", "sequence": R},
        "probe": [{"name": "P-genus", "sequence": P, "reporter": "VIC"},
                  {"name": "P-species", "sequence": species, "reporter": "FAM"}],
        "target": {"taxid": 444}, "reference_amplicon": AMP,
        "channels": [{"name": "genus", "probes": ["P-genus"]},
                     {"name": "species", "probes": ["P-species"], "target_taxid": 446}],
    })  # fmt: skip
    silent = _swap(AMP, [24, 28, 32, 36])  # four changes in the species probe site
    genome = lambda amp, s: {f"c{s}": filler(500, s) + amp + filler(500, s + 1)}  # noqa: E731
    fake = FakeDatasets([
        FakeAssembly("GCF_000000001.1", "2025-01-01", genome(AMP, 1), taxid=446),
        FakeAssembly("GCF_000000002.1", "2025-02-01", genome(silent, 3), taxid=450),
        FakeAssembly("GCF_000000003.1", "2025-03-01", genome(AMP, 5), taxid=450),
    ])  # fmt: skip
    cfg, _f, client, _a = setup(tmp_path, fake=fake)
    lineage = {446: {446, 445, 444}, 450: {450, 445, 444}}
    res = run_exhaustive(assay, cfg, client, tmp_path / "cache", _no_fetch, now=NOW,
                         ancestors_of=lambda ids: {t: lineage[t] for t in ids})  # fmt: skip
    return assay, cfg, res


def test_a_species_channel_is_judged_on_its_own_target_and_silent_elsewhere(tmp_path):
    _assay, _cfg, res = _two_channels(tmp_path)
    genus, spec = res.coverage.channel_results
    assert (genus.target_genomes, genus.detected, genus.nontarget_genomes) == (3, 3, 0)
    assert (spec.target_genomes, spec.detected) == (1, 1)
    assert (spec.nontarget_genomes, spec.silent, spec.signal) == (2, 1, 1)
    assert spec.signal_examples == ["GCF_000000003.1"]  # the species probe also binds there


def test_the_report_and_workbook_show_each_channel_and_its_status(tmp_path):
    from openpyxl import load_workbook

    from qpcr_assay_check.pipeline import evaluate
    from qpcr_assay_check.report.html import render_report
    from qpcr_assay_check.report.summary import summary_rows
    from qpcr_assay_check.report.xlsx import write_workbook
    from qpcr_assay_check.verdict import Verdict

    from .test_variants_exhaustive import _empty_specificity

    assay, cfg, res = _two_channels(tmp_path)
    result = evaluate(
        assay, cfg, now=NOW, target_sites=res.sites, variant_coverage=res.coverage,
        release_dates=res.release_dates, inclusivity=res.inclusivity,
        specificity=_empty_specificity(),
    )  # fmt: skip
    # three genomes are too few for the pooled figure (Incomplete, which outranks Review); the
    # species channel's signal in a non-target genome is added to the rationale
    assert result.inclusivity.verdict is Verdict.INCOMPLETE
    assert any(line.startswith("Channel species: a signal in 1 of 2") for line in
               result.inclusivity.rationale)  # fmt: skip
    rows = {r.check: r for r in summary_rows(result, cfg, [])}
    assert rows["Detection per channel: genus (VIC)"].css == "PASS"
    assert rows["Detection per channel: species (FAM)"].css == "WARN"
    html = render_report(result, cfg)
    assert 'id="channels"' in html and "Detection per channel" in html
    write_workbook(result, tmp_path / "r.xlsx", cfg)
    sheet = list(load_workbook(tmp_path / "r.xlsx")["Channels"].iter_rows(values_only=True))
    assert [r[0] for r in sheet[1:]] == ["genus", "species"] and sheet[2][11] == 1


def test_genomes_without_a_lineage_are_not_counted_for_a_narrower_channel(tmp_path):
    from qpcr_assay_check.models import Assay

    assay = Assay.model_validate(
        make_assay(reference_amplicon=AMP, target={"taxid": 444}).model_dump()
        | {"channels": [{"name": "species", "probes": ["probe"], "target_taxid": 446}]}
    )
    cfg, _f, client, _a = setup(tmp_path)
    res = run_exhaustive(assay, cfg, client, tmp_path / "cache", _no_fetch, now=NOW)
    (ch,) = res.coverage.channel_results
    assert ch.target_genomes == 0 and ch.membership_unknown == res.coverage.assessed_total


def test_channel_status_rules():
    from qpcr_assay_check.config import load_config
    from qpcr_assay_check.variants.exhaustive import channel_verdict
    from qpcr_assay_check.variants.models import ChannelResult
    from qpcr_assay_check.verdict import Verdict

    rules = load_config().inclusivity  # review below warn_below_percent, FAIL below fail_below
    ch = lambda **kw: ChannelResult(name="c", probes=["p"], **kw)  # noqa: E731
    ok = ch(target_genomes=100, detected=100)
    assert channel_verdict(ok, rules) == (Verdict.PASS, "")
    low = int(rules.fail_below_percent) - 1
    assert channel_verdict(ch(target_genomes=100, detected=low, not_detected=100 - low),
                           rules)[0] is Verdict.FAIL  # fmt: skip
    level, why = channel_verdict(ch(target_genomes=100, detected=100, nontarget_genomes=5,
                                    signal=2), rules)  # fmt: skip
    assert level is Verdict.WARN and "2 of 5 genomes outside its target" in why
    assert channel_verdict(ch(), rules)[0] is Verdict.INCOMPLETE
    unknown = ch(target_genomes=10, detected=10, membership_unknown=3)
    assert channel_verdict(unknown, rules)[0] is Verdict.INCOMPLETE
