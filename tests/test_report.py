import re
from datetime import UTC, datetime
from html.parser import HTMLParser

from openpyxl import load_workbook

from qpcr_assay_check.config import load_config
from qpcr_assay_check.pipeline import evaluate
from qpcr_assay_check.report.html import render_report
from qpcr_assay_check.report.xlsx import write_workbook

from .conftest import make_assay

NOW = datetime(2026, 9, 20, 12, 0, 0, tzinfo=UTC)


def render(assay, charts=False, qc_only=True):
    cfg = load_config()
    cfg.report.include_charts = charts
    return render_report(evaluate(assay, cfg, qc_only=qc_only, now=NOW), cfg)


def test_report_states_the_required_disclaimers(n1):
    html = render(n1)
    assert "does not replace experimental validation" in html
    assert "responsible for verifying this software within its own quality system" in html
    assert "no data was sent to NCBI" in html


def test_report_is_a_timestamped_version_stamped_record(n1):
    html = render(n1)
    assert "2026-09-20T12:00:00Z" in html
    assert "qpcr-assay-check v" in html
    assert re.search(r"[0-9a-f]{64}", html)  # full inputs hash


def test_full_run_report_says_incomplete_is_not_a_pass(n1):
    html = render(n1, qc_only=False)
    assert 'class="word">INCOMPLETE<' in html
    assert "not a pass" in html
    assert "Not evaluated" in html


def test_user_supplied_text_is_html_escaped():
    html = render(
        make_assay(assay_name="<script>alert(1)</script> & co", notes="<img src=x onerror=1>")
    )
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<img src=x" not in html


class _RefCollector(HTMLParser):
    """Collect attributes of real HTML tags that would make the browser fetch something."""

    def __init__(self) -> None:
        super().__init__()
        self.refs: list[tuple[str, str, str]] = []

    def handle_starttag(self, tag, attrs):
        for key, value in attrs:
            if key in {"src", "href", "action", "data", "poster", "srcset"} and value:
                self.refs.append((tag, key, value))


def test_report_with_charts_is_self_contained(n1):
    """No tag may make the browser load another file or host (links to NCBI pages are allowed).
    The Tm chart is inline SVG: no script at all, and the report stays small (it was about 5 MB
    with the inlined Plotly bundle; user, 2026-09-25)."""
    html = render(n1, charts=True)
    assert '<svg class="chart"' in html and "<script" not in html
    assert len(html.encode()) < 500_000
    parser = _RefCollector()
    parser.feed(html)
    # only plain links to NCBI record pages, which load nothing until clicked
    assert all(
        tag == "a" and key == "href" and value.startswith("https://www.ncbi.nlm.nih.gov/")
        for tag, key, value in parser.refs
    ), parser.refs
    assert "<link " not in html


def test_three_prime_end_is_marked(n1):
    assert '<b class="tail">' in render(n1)


def test_degenerate_and_modified_probe_are_explained():
    html = render(make_assay(forward="GACCCCAAAATCAGCGAAAW", probe_modifications=["MGB"]))
    assert "2 variants" in html
    assert "Tm reliability" in html


def test_workbook_has_expected_sheets(n1, tmp_path):
    cfg = load_config()
    result = evaluate(n1, cfg, qc_only=True, now=NOW)
    path = tmp_path / "r.xlsx"
    write_workbook(result, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Inputs", "Oligo QC", "Structures", "Sections"]
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Assay"] == "CDC N1"
    assert summary["Overall verdict"] == "WARN"
    assert wb["Oligo QC"].max_row == len(result.oligo_qc.checks) + 1


def test_oligo_quality_control_is_collapsed_just_before_methods(n1):
    """User 2026-09-25: the oligo checks do not change between runs and matter mainly when
    designing the PCR, so they sit folded near the end."""
    html = render(n1)
    qc = html.index("<h2>Oligo quality control</h2>")
    assert qc < html.index("<h2>Methods</h2>") and html.index("<h2>Assay as evaluated</h2>") < qc
    assert '<details class="qc">' in html[qc:] and "<h3>Hairpins and dimers</h3>" in html
    assert "<h2>Hairpins and dimers</h2>" not in html
