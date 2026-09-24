"""Variant-summary lumping (per oligo and per whole fragment), tested on directly constructed
sites -- mirrors tests/test_pairing.py's style. The end-to-end path from a target-tier search
(assess_target_sites) is tested at the bottom, on a constructed world."""

from qpcr_assay_check.specificity.models import SiteResult
from qpcr_assay_check.specificity.variants import build_variant_summary

from .conftest import make_assay

ASSAY = make_assay()
_n = 0


def mk_site(
    role,
    *,
    s_aln="AAAA",
    q_aln="AAAA",
    midline="||||",
    source="blast_full",
    tier="target",
    acc="X.1",
    org="Org",
    level="critical",
    mm=0,
    gap=0,
):
    global _n
    _n += 1
    return SiteResult(
        id=f"S{_n}", tier=tier, query=role, role=role, oligo=q_aln.replace("-", ""),
        accession=acc, taxid=1, organism=org, title="genomic", orientation="+",
        subject_start=1, subject_end=len(s_aln), source=source, q_aln=q_aln, s_aln=s_aln,
        midline=midline, n_match=len(s_aln) - mm - gap, n_mismatch=mm, n_gap=gap,
        n_ambiguous=0, defect_positions=[], mismatches_last5=0, mismatches_last3=0,
        terminal_defect=False, clean_3prime_nt=10, level=level,
    )  # fmt: skip


def test_identical_sites_lump_into_one_variant_at_100_percent():
    sites = [mk_site("forward"), mk_site("forward"), mk_site("forward")]
    summary = build_variant_summary(sites, ASSAY)
    fwd = next(o for o in summary.oligos if o.role == "forward")
    assert fwd.total_measured == 3
    assert len(fwd.rows) == 1
    assert fwd.rows[0].count == 3
    assert fwd.rows[0].percent == 100.0


def test_a_different_alignment_is_a_separate_row_with_its_own_fraction():
    sites = [
        mk_site("forward", s_aln="AAAA", midline="||||", mm=0),
        mk_site("forward", s_aln="AAAA", midline="||||", mm=0),
        mk_site("forward", s_aln="AAAT", midline="||| ", mm=1),
    ]
    summary = build_variant_summary(sites, ASSAY)
    fwd = next(o for o in summary.oligos if o.role == "forward")
    assert fwd.total_measured == 3
    assert [r.count for r in fwd.rows] == [2, 1]  # sorted, most common first
    assert [r.percent for r in fwd.rows] == [200 / 3, 100 / 3]


def test_worst_case_sites_are_excluded_from_the_table_and_counted_separately():
    sites = [
        mk_site("probe", source="blast_full"),
        mk_site("probe", source="blast_partial_worst_case"),
    ]
    summary = build_variant_summary(sites, ASSAY)
    probe = next(o for o in summary.oligos if o.role == "probe")
    assert probe.total_measured == 1
    assert probe.n_excluded_unmeasured == 1
    assert sum(r.count for r in probe.rows) == 1


def test_non_target_tier_sites_are_ignored():
    sites = [mk_site("forward", tier="background"), mk_site("forward", tier="near_neighbours")]
    summary = build_variant_summary(sites, ASSAY)
    fwd = next(o for o in summary.oligos if o.role == "forward")
    assert fwd.total_measured == 0
    assert fwd.rows == []


def triple(acc, *, mm_probe=0, probe_source="blast_full", tier="target", levels=None):
    lf, lp, lr = levels or ("critical", "critical", "critical")
    return [
        mk_site("forward", acc=acc, tier=tier, level=lf),
        mk_site("probe", acc=acc, tier=tier, s_aln="AAAA" if mm_probe == 0 else "AAAT",
                mm=mm_probe, source=probe_source, level=lp),
        mk_site("reverse", acc=acc, tier=tier, level=lr),
    ]  # fmt: skip


def test_a_complete_fragment_is_lumped_across_all_three_oligos():
    f, p, r = triple("A.1")
    summary = build_variant_summary([f, p, r], ASSAY)
    assert summary.fragment_total == 1
    assert len(summary.fragments) == 1
    row = summary.fragments[0]
    assert row.count == 1 and row.percent == 100.0
    assert row.forward.s_aln == f.s_aln and row.reverse.s_aln == r.s_aln


def test_a_record_without_a_probe_site_is_excluded_from_the_fragment_table():
    sites = [mk_site("forward", acc="A.1"), mk_site("reverse", acc="A.1")]
    summary = build_variant_summary(sites, ASSAY)
    assert summary.fragments == []
    assert summary.fragment_excluded_unmeasured == 1


def test_a_fragment_with_one_worst_case_site_is_excluded():
    summary = build_variant_summary(triple("A.1", probe_source="blast_partial_worst_case"), ASSAY)
    assert summary.fragments == []
    assert summary.fragment_excluded_unmeasured == 1


def test_two_records_with_the_same_three_variants_lump_together():
    sites = triple("A.1") + triple("B.1") + triple("C.1", mm_probe=1)
    summary = build_variant_summary(sites, ASSAY)
    assert summary.fragment_total == 3
    assert [f.count for f in summary.fragments] == [2, 1]


def test_sites_on_different_records_never_form_one_fragment():
    sites = [mk_site("forward", acc="A.1"), mk_site("probe", acc="B.1"),
             mk_site("reverse", acc="C.1")]  # fmt: skip
    summary = build_variant_summary(sites, ASSAY)
    assert summary.fragments == [] and summary.fragment_excluded_unmeasured == 3


def test_a_3prime_mismatch_primer_still_forms_a_fragment():
    """Unlike product prediction, a fragment does not require the primers to be able to prime."""
    f, p, r = triple("A.1")
    f = f.model_copy(update={"s_aln": "AAAT", "n_mismatch": 1, "terminal_defect": True,
                             "clean_3prime_nt": 0})  # fmt: skip
    summary = build_variant_summary([f, p, r], ASSAY)
    assert summary.fragment_total == 1 and summary.fragments[0].forward.s_aln == "AAAT"


def test_fragment_level_is_the_worst_of_the_three_constituent_sites():
    sites = triple("A.1", levels=("critical", "minor", "warning"))
    summary = build_variant_summary(sites, ASSAY)
    assert summary.fragments[0].level == "critical"


def test_non_target_tier_sites_never_form_a_fragment():
    summary = build_variant_summary(triple("A.1", tier="background"), ASSAY)
    assert summary.fragment_total == 0
    assert summary.fragments == []


# ------------------------------------------------ end to end: target-tier search -> variant table
def test_target_tier_hits_are_assessed_and_a_trimmed_3prime_variant_is_re_aligned(tmp_path):
    from qpcr_assay_check.config import load_config
    from qpcr_assay_check.oligo import iupac
    from qpcr_assay_check.search.orchestrate import run_search
    from qpcr_assay_check.search.planner import plan_searches
    from qpcr_assay_check.specificity.variants import assess_target_sites

    from .conftest import CDC_N1_F as F
    from .conftest import CDC_N1_P as P
    from .conftest import CDC_N1_R as R
    from .world import World, WorldFake, make_runner, mutate

    target, ref, var = 2697049, "NC_045512.2", "OT000099.1"
    w = World()
    for acc, fwd in ((ref, F), (var, mutate(F, [20]))):  # var: 3'-terminal forward mismatch
        seq = "T" * 50 + fwd + "T" * 30 + P + "T" * 30 + iupac.reverse_complement(R) + "T" * 50
        w.genome(acc, target, "SARS-CoV-2", seq)
        w.hit(target, "forward", F, acc, 51, "+", trim3=1 if acc == var else 0)
        w.hit(target, "probe", P, acc, 51 + len(F) + 30, "+")
        w.hit(target, "reverse", R, acc, 51 + len(F) + 30 + len(P) + 30, "-")

    cfg = load_config()
    cfg.search.background_taxids = []
    assay = make_assay(target={"taxid": target, "accession": ref})
    runner, store, fetcher = make_runner(cfg, tmp_path, WorldFake(w))
    plan = plan_searches(assay, cfg)
    parsed: dict = {}
    run_search(
        plan, cfg, runner, store, tmp_path / "s", inputs_hash="h", keep=parsed,
        keep_tiers={"target"},
    )  # fmt: skip

    sites = assess_target_sites(assay, cfg, plan, parsed, fetcher)
    assert len(sites) == 6  # one per record and oligo
    trimmed = next(s for s in sites if s.accession == var and s.role == "forward")
    assert trimmed.source == "realigned" and trimmed.n_mismatch == 1 and trimmed.terminal_defect

    summary = build_variant_summary(sites, assay)
    fwd = next(o for o in summary.oligos if o.role == "forward")
    assert fwd.total_measured == 2 and [r.count for r in fwd.rows] == [1, 1]
    assert summary.fragment_total == 2 and summary.fragment_excluded_unmeasured == 0


# ------------------------------------------------------------------ off-target grouping (report)
def test_off_target_sites_with_the_same_alignment_are_one_variant_across_organisms():
    from qpcr_assay_check.specificity.variants import group_off_target_sites

    near = {"s_aln": "AAAT", "midline": "||| ", "mm": 1, "level": "warning"}
    sites = [
        mk_site("forward", tier="exclusivity", acc="A.1", org="Influenza A virus", **near),
        mk_site("forward", tier="exclusivity", acc="A.1", org="Influenza A virus", **near),
        mk_site("forward", tier="exclusivity", acc="B.1", org="Influenza A virus", **near),
        mk_site("forward", tier="background", acc="C.1", org="Homo sapiens", **near),
        mk_site(
            "forward",
            tier="exclusivity",
            s_aln="ATTT",
            midline="|   ",
            mm=3,
            level="minor",
            org="Influenza A virus",
        ),  # fmt: skip
        mk_site("probe", tier="exclusivity", org="Influenza A virus", **near),  # other oligo
    ]
    groups = group_off_target_sites(sites)
    assert [(g.query, g.n_mismatch, g.n_sites) for g in groups] == [
        ("forward", 1, 4), ("probe", 1, 1), ("forward", 3, 1),
    ]  # fmt: skip
    first = groups[0]
    assert first.n_records == 3 and first.tiers == ["background", "exclusivity"]
    assert first.organisms == [("Influenza A virus", 3), ("Homo sapiens", 1)]


def test_the_report_and_workbook_show_off_target_variants_not_every_site(tmp_path):
    from openpyxl import load_workbook

    from qpcr_assay_check.config import load_config
    from qpcr_assay_check.pipeline import evaluate
    from qpcr_assay_check.report.html import render_report
    from qpcr_assay_check.report.xlsx import write_workbook

    from .test_report_exclusivity import _specificity_with_exclusivity

    near = {"s_aln": "AAAT", "midline": "||| ", "mm": 1, "level": "warning"}
    sites = [mk_site("forward", tier="exclusivity", acc=f"A{i}.1", org="Influenza A virus",
                     **near) for i in range(30)]  # fmt: skip
    cfg = load_config()
    spec = _specificity_with_exclusivity().model_copy(update={"sites": sites})
    result = evaluate(ASSAY, cfg, specificity=spec)
    html = render_report(result, cfg)
    assert html.count("A0.1") == 1 and "A29.1" not in html  # one row, not 30
    assert "Influenza A virus" in html and ">30<" in html
    write_workbook(result, tmp_path / "r.xlsx")
    wb = load_workbook(tmp_path / "r.xlsx")
    assert wb["Off-target variants"].max_row == 2 and wb["Off-target sites"].max_row == 31
