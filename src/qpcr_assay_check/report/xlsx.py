"""Excel workbook output (openpyxl)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..history.models import HistoryResult
from ..results import RunResult
from ..specificity.variants import LIST_FULL_NOTE, group_off_target_sites
from .ncbi_links import accession_url, taxon_url

_FILL = {
    "PASS": "D7EFE3",
    "WARN": "FBE7C0",
    "FAIL": "F4CFCF",
    "INFO": "E6EAEE",
    "INCOMPLETE": "D9DFE8",
}


def _sheet(
    wb: Workbook, title: str, header: list[str], rows: list[list[object]], status_col: int | None
) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    if status_col is not None:
        for r in ws.iter_rows(min_row=2):
            cell = r[status_col]
            colour = _FILL.get(str(cell.value))
            if colour:
                cell.fill = PatternFill("solid", fgColor=colour)
    for i, col in enumerate(ws.columns, start=1):
        width = min(80, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2)
        ws.column_dimensions[get_column_letter(i)].width = max(10, width)
    taxid_cols = {i for i, h in enumerate(header) if h in ("Taxonomy ID", "Taxid")}
    for r in ws.iter_rows(min_row=2):
        for i, cell in enumerate(r):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.value is None or cell.value == "":
                continue
            url = taxon_url(cell.value) if i in taxid_cols else None
            if url is None and isinstance(cell.value, str):
                url = accession_url(cell.value)  # a cell holding exactly one accession
            if url:
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"


def _history_rows(hist: HistoryResult) -> list[list[object]]:
    rows: list[list[object]] = [
        ["run", "", "Previous run", "", hist.previous_run_id or ""],
        ["run", "", "Previous generated (UTC)", "", hist.previous_generated_at or ""],
        ["run", "", "Assay/config changed since then", "", "yes" if hist.inputs_changed else "no"],
    ]
    for sc in hist.section_changes:
        if sc.changed:
            before = sc.verdict_before.value if sc.verdict_before else ""
            after = sc.verdict_after.value if sc.verdict_after else ""
            rows.append(["section", sc.title, "", before, after])
    for s in hist.new_sites:
        where = f"{s.role} {s.accession}:{s.subject_start}-{s.subject_end}"
        rows.append(["site: new", s.tier, where, "", s.level_after])
    for s in hist.resolved_sites:
        where = f"{s.role} {s.accession}:{s.subject_start}-{s.subject_end}"
        rows.append(["site: resolved", s.tier, where, s.level_before, ""])
    for s in hist.changed_sites:
        where = f"{s.role} {s.accession}:{s.subject_start}-{s.subject_end}"
        before = f"{s.level_before}, {s.n_mismatch_before} mm"
        after = f"{s.level_after}, {s.n_mismatch_after} mm"
        rows.append(["site: changed", s.tier, where, before, after])
    for a in hist.new_amplicons:
        where = f"{a.accession}:{a.start}-{a.end} ({a.roles})"
        rows.append(["amplicon: new", a.tier, where, "", a.classification or ""])
    for a in hist.resolved_amplicons:
        where = f"{a.accession}:{a.start}-{a.end} ({a.roles})"
        rows.append(["amplicon: resolved", a.tier, where, a.classification or "", ""])
    for c in hist.inclusivity_changes:
        before = (
            f"{c.percent_before:.0f}% of {c.sample_size_before}"
            if c.percent_before is not None
            else ""
        )
        after = (
            f"{c.percent_after:.0f}% of {c.sample_size_after}"
            if c.percent_after is not None
            else ""
        )
        rows.append(["inclusivity", c.role, f"year {c.year}", before, after])
    return rows


def write_workbook(result: RunResult, path: Path) -> None:
    """Write summary, inputs, QC checks, structures, and section states."""
    wb = Workbook()
    wb.remove(wb.active)
    a = result.assay
    _sheet(
        wb,
        "Summary",
        ["Item", "Value"],
        [
            ["Assay", a.assay_name],
            ["Overall verdict", result.overall.verdict.value],
            ["Mode", result.mode],
            ["Run ID", result.run_id],
            ["Generated (UTC)", result.generated_at],
            ["Tool version", result.tool["version"]],
            ["Inputs hash (SHA-256)", result.inputs_hash],
            ["Data sent to NCBI", "yes" if result.network_used else "no"],
            *(
                [["Exclusivity list source", result.exclusivity.source or "not determined"]]
                if result.exclusivity is not None
                else []
            ),
            *[["Rationale", line] for line in result.overall.rationale],
            *(
                [["Variant tables", LIST_FULL_NOTE]]
                if result.variant_summary is not None and result.variant_summary.target_list_full
                else []
            ),
        ],
        None,
    )
    _sheet(
        wb,
        "Inputs",
        ["Field", "Value"],
        [
            *[
                [
                    f"{o.role} {o.name} (5'-3')" if o.name != o.role else f"{o.role} (5'-3')",
                    o.sequence,
                ]
                for o in a.oligo_list
            ],
            *[
                [
                    f"probe {o.name} reporter / quencher / modifications"
                    if o.name != "probe"
                    else "probe reporter / quencher / modifications",
                    f"{o.reporter or ''} / {o.quencher or ''} / {', '.join(o.modifications)}",
                ]
                for o in a.probe
            ],
            ["template type", a.template_type.value],
            ["target taxid", a.target.taxid or ""],
            [
                "target excluding taxids (must not detect)",
                ", ".join(map(str, a.target.must_not_detect_taxids)),
            ],
            [
                "target excluding taxids (out of scope)",
                ", ".join(map(str, a.target.out_of_scope_taxids)),
            ],
            ["target accession", a.target.accession or ""],
            ["target gene", a.target.gene or ""],
            ["oligo source", a.oligo_source or ""],
        ],
        None,
    )
    _sheet(
        wb,
        "Oligo QC",
        ["Subject", "Check", "Value", "Unit", "Status", "Message", "Rule"],
        [
            [c.subject, c.name, c.display, c.unit, c.status.value, c.message, c.rule]
            for c in result.oligo_qc.checks
        ],
        4,
    )
    _sheet(
        wb,
        "Structures",
        [
            "Structure",
            "Kind",
            "Found",
            "Tm (°C)",
            "ΔG (kcal/mol)",
            "At (°C)",
            "Status",
            "Combinations evaluated",
        ],
        [
            [
                s.label,
                s.kind,
                "yes" if s.found else "no",
                None if s.tm_c is None else round(s.tm_c, 1),
                None if s.dg_kcal is None else round(s.dg_kcal, 2),
                s.temp_c,
                s.status.value,
                f"{s.n_evaluated}/{s.n_total}",
            ]
            for s in result.oligo_qc.structures
        ],
        6,
    )
    spec = result.specificity
    if spec is not None:
        _sheet(
            wb,
            "Off-target variants",
            ["Level", "Query", "Source", "Mismatches", "Gaps", "Clean 3' nt", "Duplex Tm (°C)",
             "ΔTm (°C)", "Sites", "Records", "Tiers", "Organisms (sites)", "Example accession",
             "Oligo", "Subject"],
            [
                [g.level, g.query, g.source, g.n_mismatch, g.n_gap, g.clean_3prime_nt,
                 None if g.tm_c is None else round(g.tm_c, 1),
                 None if g.delta_tm_c is None else round(g.delta_tm_c, 1), g.n_sites,
                 g.n_records, ", ".join(g.tiers),
                 "; ".join(f"{name} ({n})" for name, n in g.organisms),
                 g.example_site.accession, g.q_aln, g.s_aln]
                for g in group_off_target_sites(spec.sites)
            ],
            None,
        )  # fmt: skip
        _sheet(
            wb,
            "Off-target sites",
            ["Tier", "Query", "Accession", "Organism", "Strand", "Start", "End", "Level", "Source",
             "Mismatches", "Gaps", "Clean 3' nt", "Duplex Tm (°C)", "ΔTm (°C)", "Oligo", "Subject"],
            [
                [s.tier, s.query, s.accession, s.organism or "", s.orientation, s.subject_start,
                 s.subject_end, s.level, s.source, s.n_mismatch, s.n_gap, s.clean_3prime_nt,
                 None if s.tm_c is None else round(s.tm_c, 1),
                 None if s.delta_tm_c is None else round(s.delta_tm_c, 1), s.q_aln, s.s_aln]
                for s in spec.sites
            ],
            None,
        )  # fmt: skip
        _sheet(
            wb,
            "Predicted products",
            ["ID", "Tier", "Accession", "Organism", "Start", "End", "Length (bp)", "Primers",
             "Class", "Probe site", "Record", "Note"],
            [
                [a.id, a.tier, a.accession, a.organism or "", a.start, a.end, a.length, a.roles,
                 a.classification, a.probe_site or "", a.record_type, a.note]
                for a in spec.amplicons
            ],
            None,
        )  # fmt: skip
        _sheet(
            wb,
            "Searches",
            ["Tier", "Search", "Taxa", "RID", "Hits (per query)", "Perfect full-length hits"],
            [
                [r["tier"], r["label"], ", ".join(map(str, r["taxids"])), r.get("rid") or "",
                 str(r["n_hits"]), str(r.get("perfect_full_length", ""))]
                for r in spec.searches
            ],
            None,
        )  # fmt: skip
        _sheet(
            wb,
            "Findings",
            ["Severity", "Topic", "Message"],
            [[f.severity, f.topic, f.message] for f in spec.findings],
            0,
        )
    excl = result.exclusivity
    if excl is not None:
        _sheet(
            wb,
            "Exclusivity",
            ["Organism", "Resolution", "Taxonomy ID", "Sites", "Best site", "Best level",
             "Predicted product"],
            [
                [row.organism,
                 "target (excluded)" if row.is_target else row.resolution,
                 row.taxid or "", row.n_sites, row.best_site_id or "", row.best_site_level or "",
                 row.amplicon_classification or ("no" if not row.amplicon_predicted else "")]
                for row in excl.rows
            ],
            5,
        )  # fmt: skip
    if result.taxonomy_breakdown:
        _sheet(
            wb,
            "Taxonomy breakdown",
            ["Taxonomy ID", "Scientific name", "Species", "Genus", "Family", "Sites"],
            [
                [t.taxid, t.scientific_name, t.species or "", t.genus or "", t.family or "",
                 t.n_sites]
                for t in result.taxonomy_breakdown
            ],
            None,
        )  # fmt: skip
    vs = result.variant_summary
    if vs is not None:
        _sheet(
            wb,
            "Oligo variants",
            ["Oligo", "Variant (subject, aligned)", "Count", "Fraction (%)", "Mismatches",
             "Gaps", "Matching 3' nt", "Example accession", "Example organism",
             "First release", "Last release"],
            [
                [o.role if row.oligo_name in ("", o.role) else f"{o.role} {row.oligo_name}",
                 row.s_aln, row.count, round(row.percent, 2), row.n_mismatch, row.n_gap,
                 row.clean_3prime_nt, row.example_accession, row.example_organism or "",
                 row.first_seen or "", row.last_seen or ""]
                for o in vs.oligos
                for row in o.rows
            ],
            None,
        )  # fmt: skip
        _sheet(
            wb,
            "Fragment variants",
            ["Forward", "Probe", "Reverse", "Count", "Fraction (%)", "Mismatches (F/P/R)",
             "Example accession", "Example organism", "First release", "Last release"],
            [
                [f.forward.s_aln, f.probe.s_aln, f.reverse.s_aln, f.count, round(f.percent, 2),
                 f"{f.forward.n_mismatch + f.forward.n_gap}/{f.probe.n_mismatch + f.probe.n_gap}/"
                 f"{f.reverse.n_mismatch + f.reverse.n_gap}",
                 f.example_accession, f.example_organism or "",
                 f.forward.first_seen or "", f.forward.last_seen or ""]
                for f in vs.fragments
            ],
            None,
        )  # fmt: skip
        if vs.coverage is not None:
            c = vs.coverage
            _sheet(
                wb,
                "Variant coverage",
                ["Release year",
                 "Records listed" if c.source == "blast_partitioned" else "Assemblies listed",
                 "Assessed"],
                [[y.year, y.listed, y.assessed] for y in c.years]
                + [["Total", c.listed_total, c.assessed_total],
                   ["Region found (all 3 sites)", c.found, ""],
                   ["Region cut by a record end" if c.source == "blast_partitioned"
                    else "Region cut by a contig end", c.contig_break, ""],
                   ["Region hidden by N", c.masked, ", ".join(c.masked_examples)],
                   ["Region not found", c.not_found, ", ".join(c.not_found_examples)],
                   *([["  ...no sequence labelled as a plasmid", c.not_found_without_plasmid,
                       ""],
                      ["  ...plasmid sequence present, region missing (review)",
                       c.not_found_with_plasmid, ", ".join(c.not_found_with_plasmid_examples)]]
                     if c.target_on_plasmid else []),
                   ["More than one copy", c.multi_copy, ""]],
                None,
            )  # fmt: skip
            cc = c.copies
            if cc is not None and cc.genomes:
                _sheet(
                    wb,
                    "Copies and coverage",
                    ["Item", "Oligo / reporter", "Genomes", "Only this oligo", "Examples"],
                    [["Genomes assessed", "", cc.genomes, "", ""],
                     ["More than one copy", f"at most {cc.max_copies}", cc.multi_copy, "", ""],
                     ["Best copy not the first found", "", cc.best_copy_not_first, "", ""],
                     ["With a detectable copy",
                      "bulges detectable" if cc.homopolymer_bulges_detectable else "strict",
                      cc.with_detectable_copy, "", ""],
                     ["  ...strict (homopolymer bulges not detectable)", "",
                      cc.with_detectable_copy_strict, "", ""],
                     ["  ...homopolymer bulges tolerated", "", cc.with_detectable_copy_bulges,
                      "", ""],
                     ["Escapes (no detectable copy)", "", cc.escapes, "",
                      ", ".join(cc.escape_examples)],
                     *[[f"Covers ({o.role})", o.name + (f" {o.reporter}" if o.reporter else ""),
                        o.covered, o.only, ""] for o in cc.oligos],
                     *[[f"None of the {role} oligos", "", n, "",
                        ", ".join(cc.role_none_examples.get(role, []))]
                       for role, n in cc.role_none.items()],
                     *[["Channel", ch.reporter + ": " + ", ".join(ch.probes), ch.covered, "", ""]
                       for ch in cc.channels],
                     ["Any channel", cc.probe_channels, cc.any_channel, "", ""],
                     ["All channels", "", cc.all_channels, "", ""]],
                    None,
                )  # fmt: skip
    incl = result.inclusivity
    if incl is not None and incl.oligos:
        _sheet(
            wb,
            "Inclusivity",
            ["Oligo", "Year", "Population", "Sample size", "Perfect", "1 mismatch",
             "2+ mismatch/gap", "3' mismatch", "Fetch failed", "Detectable (class)",
             "At risk", "Likely failure", "Indeterminate"],
            [
                [o.role, w.year, w.population_size if w.population_size is not None else "",
                 w.sample_size, w.n_perfect, w.n_one_mismatch, w.n_two_plus_mismatch,
                 w.n_three_prime_mismatch, w.n_fetch_failed,
                 "" if w.n_detectable is None else w.n_detectable,
                 w.n_by_grade.get("at_risk", ""), w.n_by_grade.get("likely_failure", ""),
                 w.n_by_grade.get("indeterminate", "")]
                for o in incl.oligos
                for w in o.windows
            ],
            None,
        )  # fmt: skip
    hist = result.history
    if hist is not None and hist.has_previous:
        _sheet(
            wb, "History", ["Kind", "Tier/Role", "Description", "Before", "After"],
            _history_rows(hist), None,
        )  # fmt: skip
    _sheet(
        wb,
        "Sections",
        ["Section", "State", "Verdict", "Note"],
        [[s.title, s.state, s.verdict.value if s.verdict else "", s.note] for s in result.sections],
        None,
    )
    wb.save(path)
